from __future__ import annotations

import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from leadfinder.services.pipeline.quality_gate import has_verified_contact_source

EXPORT_COLUMNS = [
    "lead_id",
    "lead_score",
    "lead_tier",
    "manual_review_status",
    "do_not_contact",
    "book_title",
    "author_name",
    "illustrator_name",
    "amazon_book_url",
    "asin",
    "amazon_source_url",
    "category",
    "review_count",
    "rating",
    "publisher",
    "publication_date",
    "author_website",
    "contact_page_url",
    "public_email",
    "public_email_source_url",
    "public_email_confidence",
    "public_phone",
    "public_phone_source_url",
    "public_phone_confidence",
    "agent_name",
    "representation_email",
    "publicist_email",
    "location",
    "location_source_url",
    "instagram_url",
    "facebook_url",
    "tiktok_url",
    "youtube_url",
    "linkedin_url",
    "publisher_url",
    "video_status",
    "video_confidence",
    "video_evidence_urls",
    "video_classification_reason",
    "sales_agent_summary",
    "fit_reason",
    "suggested_pitch_angle",
    "suggested_first_line",
    "pitch_direct",
    "pitch_agent",
    "pitch_publicist",
    "what_to_say",
    "what_not_to_say",
    "next_best_action",
    "missing_data",
    "warnings",
    "all_source_urls",
    "created_at",
    "updated_at",
    "notes",
]


def _field_evidence(lead, field_name: str):
    value = getattr(lead, field_name, "")
    if not value:
        return None
    for item in lead.evidence.filter(field_name=field_name, field_value=value).order_by("-confidence", "-created_at"):
        min_confidence = 0.55 if field_name == "public_phone" else 0.6
        if has_verified_contact_source(lead, field_name, item.field_value, min_confidence=min_confidence):
            return item
    return None


def _verified_contact_value(lead, field_name: str) -> str:
    value = getattr(lead, field_name, "")
    if not value:
        return ""
    min_confidence = 0.55 if field_name == "public_phone" else 0.6
    return value if has_verified_contact_source(lead, field_name, value, min_confidence=min_confidence) else ""


def lead_to_row(lead) -> dict:
    book = lead.book
    author = lead.author_profile
    email_evidence = _field_evidence(lead, "public_email")
    phone_evidence = _field_evidence(lead, "public_phone")
    verified_email = _verified_contact_value(lead, "public_email")
    verified_phone = _verified_contact_value(lead, "public_phone")
    location_evidence = _field_evidence(lead, "location")
    videos = list(lead.videos.all())
    all_urls = list(lead.evidence.exclude(source_url="").values_list("source_url", flat=True).distinct())
    brief = getattr(lead, "brief", None)
    return {
        "lead_id": lead.id,
        "lead_score": lead.lead_score,
        "lead_tier": lead.lead_tier,
        "manual_review_status": lead.manual_review_status,
        "do_not_contact": lead.do_not_contact,
        "book_title": book.title,
        "author_name": book.author_name,
        "illustrator_name": book.illustrator_name,
        "amazon_book_url": book.amazon_book_url,
        "asin": book.asin,
        "amazon_source_url": book.amazon_source_url,
        "category": book.category,
        "review_count": book.review_count,
        "rating": book.rating,
        "publisher": book.publisher,
        "publication_date": book.publication_date,
        "author_website": author.canonical_website if author else "",
        "contact_page_url": author.contact_page_url if author else "",
        "public_email": verified_email,
        "public_email_source_url": email_evidence.source_url if email_evidence else "",
        "public_email_confidence": email_evidence.confidence if email_evidence else "",
        "public_phone": verified_phone,
        "public_phone_source_url": phone_evidence.source_url if phone_evidence else "",
        "public_phone_confidence": phone_evidence.confidence if phone_evidence else "",
        "agent_name": author.agent_name if author else "",
        "representation_email": lead.representation_email or (author.representation_email if author else ""),
        "publicist_email": lead.publicist_email or (author.publicist_email if author else ""),
        "location": lead.location,
        "location_source_url": location_evidence.source_url if location_evidence else "",
        "instagram_url": author.instagram_url if author else "",
        "facebook_url": author.facebook_url if author else "",
        "tiktok_url": author.tiktok_url if author else "",
        "youtube_url": author.youtube_url if author else "",
        "linkedin_url": author.linkedin_url if author else "",
        "publisher_url": author.publisher_url if author else "",
        "video_status": lead.video_status,
        "video_confidence": lead.video_confidence,
        "video_evidence_urls": "; ".join(video.video_url for video in videos),
        "video_classification_reason": "; ".join(filter(None, [video.classification_reason for video in videos])),
        "sales_agent_summary": lead.sales_agent_summary,
        "fit_reason": lead.fit_reason,
        "suggested_pitch_angle": lead.suggested_pitch_angle,
        "suggested_first_line": lead.suggested_first_line,
        "pitch_direct": brief.pitch_direct if brief else "",
        "pitch_agent": brief.pitch_agent if brief else "",
        "pitch_publicist": brief.pitch_publicist if brief else "",
        "what_to_say": lead.what_to_say,
        "what_not_to_say": lead.what_not_to_say,
        "next_best_action": lead.next_best_action,
        "missing_data": lead.missing_data_json,
        "warnings": lead.warnings_json,
        "all_source_urls": "; ".join(all_urls),
        "created_at": lead.created_at,
        "updated_at": lead.updated_at,
        "notes": lead.notes,
    }


def export_leads_response(queryset) -> HttpResponse:
    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = 'attachment; filename="book_trailer_leads.csv"'
    writer = csv.DictWriter(response, fieldnames=EXPORT_COLUMNS)
    writer.writeheader()
    for lead in queryset.select_related("book", "author_profile").prefetch_related("evidence", "videos"):
        writer.writerow(lead_to_row(lead))
    return response


def export_leads_to_file(queryset, output_path: str) -> int:
    count = 0
    with open(output_path, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=EXPORT_COLUMNS)
        writer.writeheader()
        for lead in queryset.select_related("book", "author_profile").prefetch_related("evidence", "videos"):
            writer.writerow(lead_to_row(lead))
            count += 1
    return count


def build_leads_workbook(queryset) -> openpyxl.Workbook:
    """
    Constructs a 4-tab Excel workbook styled professionally matching the Lead SOP.
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    # Styling Tokens
    font_family = "Segoe UI"
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    data_font = Font(name=font_family, size=10)
    
    # Premium Header Theme (Steel/Navy Teal)
    header_fill = PatternFill(start_color="1E3D59", end_color="1E3D59", fill_type="solid")
    
    # soft priority colors
    hot_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")       # Light green
    warm_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")      # Light yellow
    cold_fill = PatternFill(start_color="D1ECF1", end_color="D1ECF1", fill_type="solid")      # Light blue
    rejected_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")  # Light red
    
    thin_border = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )
    
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # ------------------
    # TAB 1: Authors
    # ------------------
    ws1 = wb.create_sheet(title="Authors")
    headers1 = [
        "Lead ID", "Date Added", "Author Name", "Pen Name", "Country", "City/State",
        "Author Type", "Primary Niche", "Official Website", "Public Email", "Contact Form",
        "Amazon Author Page", "Goodreads", "BookBub", "Instagram", "Facebook", "TikTok",
        "YouTube", "LinkedIn", "Verification Status", "Lead Score", "Priority", "Outreach Status"
    ]
    ws1.append(headers1)
    
    # ------------------
    # TAB 2: Books
    # ------------------
    ws2 = wb.create_sheet(title="Books")
    headers2 = [
        "Lead ID", "Book Title", "ASIN", "ISBN", "Amazon URL", "Publisher/Imprint",
        "Publication Date", "Format", "Review Count", "Rating", "Category", 
        "Cover Quality", "Description Quality", "A+ Content Present", "Service Need Tags", "Notes"
    ]
    ws2.append(headers2)
    
    # ------------------
    # TAB 3: Contact Evidence
    # ------------------
    ws3 = wb.create_sheet(title="Contact Evidence")
    headers3 = [
        "Lead ID", "Contact Type", "Contact Value", "Source URL", "Source Type",
        "Date Verified", "Verified By", "Confidence Level", "Notes"
    ]
    ws3.append(headers3)
    
    # ------------------
    # TAB 4: Outreach Log
    # ------------------
    ws4 = wb.create_sheet(title="Outreach Log")
    headers4 = [
        "Lead ID", "Outreach Angle", "Service Pitch", "Email/DM Sent Date",
        "Follow-Up 1", "Follow-Up 2", "Reply Status", "Call Booked", "Deal Status",
        "Notes", "Do Not Contact", "Direct Pitch", "Agent Pitch", "Publicist Pitch"
    ]
    ws4.append(headers4)

    # Style all headers
    for ws in [ws1, ws2, ws3, ws4]:
        ws.row_dimensions[1].height = 28
        for col_idx, cell in enumerate(ws[1], 1):
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

    # Fetch data
    leads = queryset.select_related("book", "author_profile").prefetch_related("evidence", "videos", "brief")
    
    for lead in leads:
        book = lead.book
        author = lead.author_profile
        brief = getattr(lead, "brief", None)
        verified_email = _verified_contact_value(lead, "public_email")
        
        # Determine values
        author_name = author.author_name if author else book.author_name
        city_state = lead.location or (author.location if author else "")
        niche = book.category
        website = author.canonical_website if author else ""
        contact_form = author.contact_page_url if author else ""
        amazon_author = author.amazon_author_url if author else ""
        goodreads = author.goodreads_url if author else ""
        
        # Verify status
        v_conf = author.identity_confidence if author else 0
        v_status = "Verified" if v_conf >= 0.8 else "Partially verified" if v_conf >= 0.5 else "Unverified"
        
        # Priority mapping
        priority_label = "Low Priority"
        if lead.lead_tier == "hot":
            priority_label = "A Lead"
        elif lead.lead_tier == "warm":
            priority_label = "B Lead"
        elif lead.lead_tier == "cold":
            priority_label = "C Lead"
            
        row1 = [
            str(lead.id),
            lead.created_at.strftime("%Y-%m-%d") if lead.created_at else "",
            author_name,
            "",  # Pen Name
            "US" if city_state else "",  # Country
            city_state,
            "KDP Self-Published Author" if book.publisher == "Independently published" else "Indie Author",
            niche,
            website,
            verified_email,
            contact_form,
            amazon_author,
            goodreads,
            "",  # BookBub
            author.instagram_url if author else "",
            author.facebook_url if author else "",
            author.tiktok_url if author else "",
            author.youtube_url if author else "",
            author.linkedin_url if author else "",
            v_status,
            lead.lead_score,
            priority_label,
            lead.manual_review_status
        ]
        ws1.append(row1)
        
        # Write to Book sheet
        cover_quality = "Weak" if book.rating and book.rating < 4.2 else "Good"
        desc_quality = "Needs Optimization" if book.amazon_source_snippet and len(book.amazon_source_snippet) < 250 else "Good"
        ap_present = "Yes" if book.has_aplus_content is True else "No" if book.has_aplus_content is False else "Unknown"
        service_tags = ", ".join(lead.service_needs_json or [])
        
        row2 = [
            str(lead.id),
            book.title,
            book.asin,
            "",  # ISBN
            book.amazon_book_url,
            book.publisher,
            book.publication_date,
            "Paperback" if "paperback" in (book.category or "").lower() else "Hardcover" if "hardcover" in (book.category or "").lower() else "Kindle Edition",
            book.review_count,
            float(book.rating) if book.rating else None,
            book.category,
            cover_quality,
            desc_quality,
            ap_present,
            service_tags,
            lead.notes
        ]
        ws2.append(row2)
        
        # Write to Evidence sheet
        evidence_list = lead.evidence.all()
        for ev in evidence_list:
            row3 = [
                str(lead.id),
                ev.field_name,
                ev.field_value,
                ev.source_url,
                ev.evidence_type,
                ev.created_at.strftime("%Y-%m-%d %H:%M") if ev.created_at else "",
                "AI Pipeline",
                ev.confidence,
                ev.source_title
            ]
            ws3.append(row3)
            
        # Write to Outreach sheet
        row4 = [
            str(lead.id),
            lead.suggested_pitch_angle,
            lead.what_to_say,
            "",  # Email Sent Date
            "",  # Follow-Up 1
            "",  # Follow-Up 2
            "",  # Reply Status
            "",  # Call Booked
            "",  # Deal Status
            lead.fit_reason,
            "Yes" if lead.do_not_contact else "No",
            brief.pitch_direct if brief else "",
            brief.pitch_agent if brief else "",
            brief.pitch_publicist if brief else ""
        ]
        ws4.append(row4)

    # Style data rows in Authors sheet
    for r_idx in range(2, ws1.max_row + 1):
        ws1.row_dimensions[r_idx].height = 20
        tier_val = ws1.cell(row=r_idx, column=22).value  # Priority is at column 22
        
        row_fill = None
        if tier_val == "A Lead":
            row_fill = hot_fill
        elif tier_val == "B Lead":
            row_fill = warm_fill
        elif tier_val == "C Lead":
            row_fill = cold_fill
        elif tier_val == "Low Priority":
            row_fill = rejected_fill
            
        for c_idx in range(1, ws1.max_column + 1):
            cell = ws1.cell(row=r_idx, column=c_idx)
            cell.font = data_font
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill
            
            # Alignments
            if c_idx in [1, 2, 4, 5, 20, 21, 22, 23]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # Style data rows in Books sheet
    for r_idx in range(2, ws2.max_row + 1):
        ws2.row_dimensions[r_idx].height = 20
        row_fill = None
        lead_id_val = ws2.cell(row=r_idx, column=1).value
        for a_idx in range(2, ws1.max_row + 1):
            if ws1.cell(row=a_idx, column=1).value == lead_id_val:
                tier_val = ws1.cell(row=a_idx, column=22).value
                if tier_val == "A Lead":
                    row_fill = hot_fill
                elif tier_val == "B Lead":
                    row_fill = warm_fill
                elif tier_val == "C Lead":
                    row_fill = cold_fill
                elif tier_val == "Low Priority":
                    row_fill = rejected_fill
                break
                
        for c_idx in range(1, ws2.max_column + 1):
            cell = ws2.cell(row=r_idx, column=c_idx)
            cell.font = data_font
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill
            if c_idx in [1, 3, 4, 7, 8, 9, 10, 12, 13, 14]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # Style data rows in Contact Evidence sheet
    for r_idx in range(2, ws3.max_row + 1):
        ws3.row_dimensions[r_idx].height = 20
        row_fill = None
        lead_id_val = ws3.cell(row=r_idx, column=1).value
        for a_idx in range(2, ws1.max_row + 1):
            if ws1.cell(row=a_idx, column=1).value == lead_id_val:
                tier_val = ws1.cell(row=a_idx, column=22).value
                if tier_val == "A Lead":
                    row_fill = hot_fill
                elif tier_val == "B Lead":
                    row_fill = warm_fill
                elif tier_val == "C Lead":
                    row_fill = cold_fill
                elif tier_val == "Low Priority":
                    row_fill = rejected_fill
                break
                
        for c_idx in range(1, ws3.max_column + 1):
            cell = ws3.cell(row=r_idx, column=c_idx)
            cell.font = data_font
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill
            if c_idx in [1, 5, 6, 7, 8]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # Style data rows in Outreach Log sheet
    for r_idx in range(2, ws4.max_row + 1):
        ws4.row_dimensions[r_idx].height = 20
        row_fill = None
        lead_id_val = ws4.cell(row=r_idx, column=1).value
        for a_idx in range(2, ws1.max_row + 1):
            if ws1.cell(row=a_idx, column=1).value == lead_id_val:
                tier_val = ws1.cell(row=a_idx, column=22).value
                if tier_val == "A Lead":
                    row_fill = hot_fill
                elif tier_val == "B Lead":
                    row_fill = warm_fill
                elif tier_val == "C Lead":
                    row_fill = cold_fill
                elif tier_val == "Low Priority":
                    row_fill = rejected_fill
                break
                
        for c_idx in range(1, ws4.max_column + 1):
            cell = ws4.cell(row=r_idx, column=c_idx)
            cell.font = data_font
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill
            if c_idx in [1, 4, 5, 6, 7, 8, 9, 11]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # Auto-fit column widths for all sheets
    for ws in [ws1, ws2, ws3, ws4]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = min(45, max(12, max_len + 3))

    return wb


def export_leads_xlsx(queryset) -> HttpResponse:
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="book_trailer_leads.xlsx"'
    wb = build_leads_workbook(queryset)
    wb.save(response)
    return response


def export_leads_xlsx_to_file(queryset, output_path: str) -> int:
    wb = build_leads_workbook(queryset)
    wb.save(output_path)
    return queryset.count()
