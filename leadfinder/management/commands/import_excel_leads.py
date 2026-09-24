from __future__ import annotations

import logging
from pathlib import Path
from urllib.parse import urlparse

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone
import openpyxl

from leadfinder.models import AuthorProfile, Book, ContactCandidate, Evidence, Lead, ResearchRun
from leadfinder.services.amazon.amazon_url_parser import extract_asin, normalize_amazon_book_url
from leadfinder.utils.normalize import normalize_text, normalized_author_key, normalized_book_key
from leadfinder.views import _status_value

logger = logging.getLogger(__name__)


def _clean_pub_date(value) -> str:
    if not value:
        return ""
    val = str(value).strip()
    if val.endswith(".0") and val[:-2].isdigit():
        return val[:-2]
    if " 00:00:00" in val:
        val = val.replace(" 00:00:00", "").strip()
    import re
    m1 = re.search(r"\b(19\d{2}|20\d{2})[-/](\d{1,2})[-/](\d{1,2})\b", val)
    if m1:
        return f"{m1.group(1)}-{int(m1.group(2)):02d}-{int(m1.group(3)):02d}"
    m_written = re.search(r"(January|February|March|April|May|June|July|August|September|October|November|December)\s+(\d{1,2}),?\s+(20\d{2}|19\d{2})", val, re.I)
    if m_written:
        months = {"january":1,"february":2,"march":3,"april":4,"may":5,"june":6,"july":7,"august":8,"september":9,"october":10,"november":11,"december":12}
        mo = months[m_written.group(1).lower()]
        return f"{m_written.group(3)}-{mo:02d}-{int(m_written.group(2)):02d}"
    m2 = re.search(r"\b(19\d{2}|20\d{2})[-/](\d{1,2})\b", val)
    if m2:
        return f"{m2.group(1)}-{int(m2.group(2)):02d}"
    m3 = re.search(r"\b(19\d{2}|20\d{2})\b", val)
    if m3:
        return m3.group(1)
    return ""


def _extract_root_website(url: str) -> str:
    if not url or not url.startswith(("http://", "https://")):
        return ""
    parsed = urlparse(url)
    if parsed.netloc:
        return f"{parsed.scheme}://{parsed.netloc}"
    return ""


def _parse_children_proof(value: str) -> tuple[bool | None, bool | None]:
    text = (value or "").strip().lower()
    if not text:
        return None, None
    if text.startswith("no"):
        return False, False
    is_children = None
    is_picture = None
    if any(k in text for k in ("children", "kid", "picture", "illustrated", "illustration", "yes")):
        is_children = "children" in text or "kid" in text or text == "yes"
        is_picture = "picture" in text or "illustrated" in text or "illustration" in text
    return is_children, is_picture


def _clean_contact_value(value) -> str:
    val = str(value or "").strip()
    if val.lower() in {"", "none", "no", "n/a", "null", "false"}:
        return ""
    return val


class Command(BaseCommand):
    help = "Import and deduplicate book leads from Excel workbooks (e.g. leads data.xlsx, leads dev-ali.xlsx)."

    def add_arguments(self, parser):
        parser.add_argument("paths", nargs="+", help="Path(s) to the Excel workbook(s) to import.")
        parser.add_argument("--run-name", default=None, help="Custom keyword / run name for the research run.")

    def handle(self, *args, **options):
        paths = [Path(p) for p in options["paths"]]
        for path in paths:
            if not path.is_file():
                raise CommandError(f"Workbook file not found: {path}")

        total_processed_all = 0
        total_created_all = 0
        total_enriched_all = 0

        for path in paths:
            self.stdout.write(self.style.NOTICE(f"\nProcessing workbook: {path}"))
            run_keyword = options.get("run_name") or f"Manual import: {path.name}"
            created_cnt, enriched_cnt, stats = self._import_file(path, run_keyword)
            total_processed_all += (created_cnt + enriched_cnt)
            total_created_all += created_cnt
            total_enriched_all += enriched_cnt
            self.stdout.write(
                self.style.SUCCESS(
                    f"Finished {path.name}: {created_cnt} books/leads created, {enriched_cnt} deduplicated/enriched.\n"
                    f"  Status breakdown: {stats['verified']} verified, {stats['other']} needs review.\n"
                    f"  Contacts: {stats['with_email']} with email, {stats['with_phone']} with phone."
                )
            )

        self.stdout.write(
            self.style.SUCCESS(
                f"\n=== All Imports Completed ==="
                f"\nTotal rows processed: {total_processed_all}"
                f"\nTotal new records created: {total_created_all}"
                f"\nTotal records deduplicated/enriched: {total_enriched_all}"
            )
        )

    def _import_file(self, path: Path, run_keyword: str) -> tuple[int, int, dict]:
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.active

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            self.stdout.write(self.style.WARNING(f"Workbook {path.name} is empty."))
            return 0, 0, {"verified": 0, "other": 0, "with_email": 0, "with_phone": 0}

        headers = [str(c).strip() if c is not None else "" for c in rows[0]]
        non_empty = [
            dict(zip(headers, r))
            for r in rows[1:]
            if any(c is not None and str(c).strip() != "" for c in r)
        ]

        # Filter out rows with no book name
        valid_rows = [r for r in non_empty if (r.get("Book Name") or "").strip()]

        created_count = 0
        enriched_count = 0
        stats = {"verified": 0, "other": 0, "with_email": 0, "with_phone": 0}

        with transaction.atomic():
            now = timezone.now()
            run = ResearchRun.objects.create(
                keyword=run_keyword,
                source_provider="manual",
                marketplace="amazon.com",
                max_books=len(valid_rows),
                status="completed",
                started_at=now,
                completed_at=now,
                settings_json={
                    "source": "excel_import",
                    "file_name": path.name,
                    "row_count": len(valid_rows),
                    "verify_imported_contacts": False,
                    "run_video_search": False,
                    "run_groq_ai_extraction": False,
                },
            )

            for r in valid_rows:
                title = str(r.get("Book Name") or "").strip()
                author_name = str(r.get("Author / Owner Name") or "").strip()
                amazon_url = str(r.get("Amazon Book / Author URL") or "").strip()
                if amazon_url.lower() in {"none", "n/a", "null"}:
                    amazon_url = ""

                asin = (extract_asin(amazon_url) or "").upper()
                if asin and not amazon_url:
                    amazon_url = normalize_amazon_book_url(f"https://www.amazon.com/dp/{asin}")
                elif amazon_url and asin:
                    amazon_url = normalize_amazon_book_url(amazon_url)

                category = str(r.get("Category") or "").strip()
                pub_date = _clean_pub_date(r.get("Publication Date"))
                children_proof = str(r.get("Children/Picture/Illustration Proof") or "").strip()
                is_children, is_picture = _parse_children_proof(children_proof)

                raw_email = _clean_contact_value(r.get("Email"))
                raw_phone = _clean_contact_value(r.get("Phone"))
                raw_dial_format = _clean_contact_value(r.get("Phone Dial Format"))
                email_proof_url = str(r.get("Email Proof URL") or "").strip()
                phone_proof_url = str(r.get("Phone Proof URL") or "").strip()
                if not raw_email and "@" in email_proof_url and not email_proof_url.startswith(("http://", "https://")):
                    raw_email = email_proof_url
                    if phone_proof_url.startswith(("http://", "https://")):
                        email_proof_url = phone_proof_url
                        phone_proof_url = ""
                    else:
                        email_proof_url = ""
                location = str(r.get("Phone Country / Region") or r.get("Phone Country/Region") or "").strip()
                rationale = str(r.get("Owner Contact Rationale") or "").strip()
                notes = str(r.get("Notes") or "").strip()
                quality_status = str(r.get("Lead Quality Status") or "").strip()
                contact_channel = str(r.get("Contact Source Channel") or "").strip()

                status = _status_value(quality_status)
                score = 90 if status == "verified" else 0

                if status == "verified":
                    stats["verified"] += 1
                else:
                    stats["other"] += 1

                if raw_email:
                    stats["with_email"] += 1
                if raw_phone:
                    stats["with_phone"] += 1

                # Clean phone dial format: ensure leading '+' if numeric international
                clean_phone_norm = raw_dial_format or raw_phone
                if clean_phone_norm and clean_phone_norm.isdigit() and len(clean_phone_norm) >= 10:
                    clean_phone_norm = f"+{clean_phone_norm}"

                # Deduplication lookup
                book_key = normalized_book_key(title, author_name, asin)
                title_author_key = normalized_book_key(title, author_name, "")
                author_key = normalized_author_key(author_name)

                existing_book = (
                    Book.objects.filter(normalized_key=book_key).first()
                    or (asin and Book.objects.filter(asin=asin).first())
                    or Book.objects.filter(normalized_key=title_author_key).first()
                )

                canonical_web = _extract_root_website(email_proof_url)
                amazon_author_url = amazon_url if ("/e/" in amazon_url or "author" in amazon_url.lower()) else ""

                # Handle AuthorProfile
                author_profile = None
                if author_name:
                    author_profile = AuthorProfile.objects.filter(normalized_author_key=author_key).first()
                    if not author_profile:
                        author_profile = AuthorProfile.objects.create(
                            author_name=author_name,
                            normalized_author_key=author_key,
                            canonical_website=canonical_web,
                            contact_page_url=email_proof_url,
                            amazon_author_url=amazon_author_url,
                            location=location,
                            identity_confidence=0.9 if status == "verified" else 0.7,
                            identity_reason=quality_status or "Author lead from client workbook import.",
                        )
                    else:
                        profile_updates = []
                        if not author_profile.canonical_website and canonical_web:
                            author_profile.canonical_website = canonical_web
                            profile_updates.append("canonical_website")
                        if not author_profile.contact_page_url and email_proof_url:
                            author_profile.contact_page_url = email_proof_url
                            profile_updates.append("contact_page_url")
                        if not author_profile.amazon_author_url and amazon_author_url:
                            author_profile.amazon_author_url = amazon_author_url
                            profile_updates.append("amazon_author_url")
                        if not author_profile.location and location:
                            author_profile.location = location
                            profile_updates.append("location")
                        if profile_updates:
                            author_profile.save(update_fields=[*profile_updates, "updated_at"])

                if existing_book:
                    enriched_count += 1
                    book = existing_book
                    book_updates = []
                    if not book.asin and asin:
                        book.asin = asin
                        book_updates.append("asin")
                    if not book.amazon_book_url and amazon_url:
                        book.amazon_book_url = amazon_url
                        book_updates.append("amazon_book_url")
                    if not book.category and category:
                        book.category = category
                        book_updates.append("category")
                    if not book.publication_date and pub_date:
                        book.publication_date = pub_date
                        book_updates.append("publication_date")
                    if book.is_childrens_book is None and is_children is not None:
                        book.is_childrens_book = is_children
                        book_updates.append("is_childrens_book")
                    if book.is_picture_or_illustrated_book is None and is_picture is not None:
                        book.is_picture_or_illustrated_book = is_picture
                        book_updates.append("is_picture_or_illustrated_book")
                    if book_updates:
                        book.save(update_fields=[*book_updates, "updated_at"])

                    lead = Lead.objects.filter(book=book).first()
                    if not lead:
                        lead = Lead.objects.create(
                            book=book,
                            author_profile=author_profile,
                            public_email=raw_email,
                            public_phone=raw_phone,
                            location=location,
                            lead_score=score,
                            verification_score=score,
                            verification_status=status,
                            verification_reason=quality_status or "Imported from Excel workbook.",
                            verification_reasons_json=["uploader_verified", "source_proof_provided"] if status == "verified" else ["uploader_attested", "import_awaiting_verification"],
                            verification_version="manual-import-v2",
                            verified_at=timezone.now() if status == "verified" else None,
                            uploader_attested=True,
                            uploader_attested_at=timezone.now(),
                            sales_agent_summary=rationale,
                            fit_reason=rationale,
                            notes=notes,
                            manual_review_status="verified" if status == "verified" else "needs_review",
                        )
                    else:
                        lead_updates = []
                        if not lead.public_email and raw_email:
                            lead.public_email = raw_email
                            lead_updates.append("public_email")
                        if not lead.public_phone and raw_phone:
                            lead.public_phone = raw_phone
                            lead_updates.append("public_phone")
                        if not lead.location and location:
                            lead.location = location
                            lead_updates.append("location")
                        if not lead.sales_agent_summary and rationale:
                            lead.sales_agent_summary = rationale
                            lead.fit_reason = rationale
                            lead_updates.extend(["sales_agent_summary", "fit_reason"])
                        if not lead.notes and notes:
                            lead.notes = notes
                            lead_updates.append("notes")
                        if status == "verified" and lead.verification_status != "verified":
                            lead.verification_status = "verified"
                            lead.verification_score = score
                            lead.manual_review_status = "verified"
                            lead.verification_reason = quality_status
                            lead.verification_reasons_json = ["uploader_verified", "source_proof_provided"]
                            lead.verified_at = timezone.now()
                            lead_updates.extend(["verification_status", "verification_score", "manual_review_status", "verification_reason", "verification_reasons_json", "verified_at"])
                        if not lead.author_profile and author_profile:
                            lead.author_profile = author_profile
                            lead_updates.append("author_profile")
                        if lead_updates:
                            lead.save(update_fields=[*lead_updates, "updated_at"])
                else:
                    created_count += 1
                    book = Book.objects.create(
                        research_run=run,
                        title=title,
                        author_name=author_name,
                        asin=asin,
                        amazon_book_url=amazon_url,
                        amazon_source_url=amazon_url or path.name,
                        category=category,
                        publication_date=pub_date,
                        normalized_key=book_key,
                        book_data_confidence=0.9,
                        source_provider="manual",
                        source_raw_json={"source": path.name},
                        is_childrens_book=is_children,
                        is_picture_or_illustrated_book=is_picture,
                    )

                    lead = Lead.objects.create(
                        book=book,
                        author_profile=author_profile,
                        public_email=raw_email,
                        public_phone=raw_phone,
                        location=location,
                        lead_score=score,
                        verification_score=score,
                        verification_status=status,
                        verification_reason=quality_status or "Imported from Excel workbook.",
                        verification_reasons_json=["uploader_verified", "source_proof_provided"] if status == "verified" else ["uploader_attested", "import_awaiting_verification"],
                        verification_version="manual-import-v2",
                        verified_at=timezone.now() if status == "verified" else None,
                        uploader_attested=True,
                        uploader_attested_at=timezone.now(),
                        sales_agent_summary=rationale,
                        fit_reason=rationale,
                        notes=notes,
                        manual_review_status="verified" if status == "verified" else "needs_review",
                    )

                # Contacts and Evidence
                has_primary = ContactCandidate.objects.filter(lead=lead, is_primary=True).exists()

                # Email Candidate
                if raw_email:
                    email_norm = raw_email.lower().strip()
                    email_cand = ContactCandidate.objects.filter(lead=lead, channel="email", normalized_value=email_norm).first()
                    if not email_cand:
                        email_cand = ContactCandidate.objects.create(
                            lead=lead,
                            channel="email",
                            role="author",
                            raw_value=raw_email,
                            normalized_value=email_norm,
                            verification_status=status,
                            verification_score=score,
                            deliverability_status="deliverable" if status == "verified" else "manual_import",
                            is_primary=not has_primary,
                            selected_reason=quality_status or ("Verified author email" if status == "verified" else "Imported email; system verification not run."),
                        )
                        if not has_primary:
                            has_primary = True
                            lead.primary_contact = email_cand
                            lead.save(update_fields=["primary_contact", "updated_at"])
                    elif status == "verified" and email_cand.verification_status != "verified":
                        email_cand.verification_status = "verified"
                        email_cand.verification_score = score
                        email_cand.deliverability_status = "deliverable"
                        email_cand.selected_reason = quality_status
                        email_cand.save(update_fields=["verification_status", "verification_score", "deliverability_status", "selected_reason", "updated_at"])
                        if not lead.primary_contact:
                            lead.primary_contact = email_cand
                            lead.save(update_fields=["primary_contact", "updated_at"])

                # Phone Candidate
                if raw_phone and clean_phone_norm:
                    phone_cand = ContactCandidate.objects.filter(lead=lead, channel="phone", normalized_value=clean_phone_norm).first()
                    if not phone_cand:
                        phone_cand = ContactCandidate.objects.create(
                            lead=lead,
                            channel="phone",
                            role="author",
                            raw_value=raw_phone,
                            normalized_value=clean_phone_norm,
                            verification_status=status,
                            verification_score=score,
                            deliverability_status="not_applicable",
                            is_primary=not has_primary,
                            selected_reason=quality_status or "Author phone contact",
                        )
                        if not has_primary:
                            has_primary = True
                            lead.primary_contact = phone_cand
                            lead.save(update_fields=["primary_contact", "updated_at"])
                    elif status == "verified" and phone_cand.verification_status != "verified":
                        phone_cand.verification_status = "verified"
                        phone_cand.verification_score = score
                        phone_cand.selected_reason = quality_status
                        phone_cand.save(update_fields=["verification_status", "verification_score", "selected_reason", "updated_at"])
                        if not lead.primary_contact:
                            lead.primary_contact = phone_cand
                            lead.save(update_fields=["primary_contact", "updated_at"])

                # Evidence records
                if amazon_url:
                    Evidence.objects.get_or_create(
                        book=book,
                        field_name="amazon_book_url",
                        defaults={
                            "lead": lead,
                            "author_profile": author_profile,
                            "evidence_type": "manual",
                            "field_value": amazon_url,
                            "source_url": path.name,
                            "source_title": "Amazon listing",
                            "confidence": 0.9,
                            "is_primary": True,
                        },
                    )

                if raw_email and email_proof_url:
                    Evidence.objects.get_or_create(
                        lead=lead,
                        field_name="public_email",
                        source_url=email_proof_url,
                        defaults={
                            "book": book,
                            "author_profile": author_profile,
                            "evidence_type": "manual",
                            "field_value": raw_email,
                            "source_title": contact_channel or "Author Website",
                            "confidence": 0.85,
                            "is_primary": True,
                        },
                    )

                if raw_phone and phone_proof_url:
                    Evidence.objects.get_or_create(
                        lead=lead,
                        field_name="public_phone",
                        source_url=phone_proof_url,
                        defaults={
                            "book": book,
                            "author_profile": author_profile,
                            "evidence_type": "manual",
                            "field_value": raw_phone,
                            "source_title": contact_channel or "Author Website",
                            "confidence": 0.85,
                            "is_primary": True,
                        },
                    )

            run.max_books = run.books.count()
            run.save(update_fields=["max_books"])

        return created_count, enriched_count, stats
