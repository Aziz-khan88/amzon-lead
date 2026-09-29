from __future__ import annotations

import csv
import io
import importlib.util
import logging
import os
import re
import threading
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from openpyxl import load_workbook
from django.conf import settings
from django.contrib import messages
from django.contrib.auth import get_user_model
from django.core.paginator import Paginator
from django.db import close_old_connections, transaction
from django.db.models import Count, Exists, OuterRef, Prefetch, Q, Subquery
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.views.decorators.http import require_GET, require_POST
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme

from .forms import (
    BookLifeRunForm,
    CSVImportForm,
    ISBNSearchForm,
    LeadFilterForm,
    LeadAssignmentScheduleForm,
    LeadAssignmentUpdateForm,
    ResearchRunForm,
    ScheduledLeadTaskForm,
    TeamMemberCreateForm,
)
from .models import (
    AuthorProfile, Book, ContactCandidate, Evidence, Lead, LeadAssignment,
    LeadAssignmentSchedule, LeadAssignmentScheduleRun, ResearchRun, ScheduledLeadTask,
    UserProfile, VerificationBatch,
)
from .access import ensure_lead_access, is_management, lead_queryset_for_user, roles_required, user_role
from .services.assignments import assign_leads, calculate_next_assignment_at, execute_assignment_schedule
from .services.amazon.amazon_url_parser import extract_asin, is_amazon_url, normalize_amazon_book_url
from .services.booklife import BOOKLIFE_CATEGORIES, grouped_booklife_categories
from .services.books.isbn_intelligence import analyze_identifier, generate_barcode_svg
from .services.export.csv_export import export_leads_response, export_leads_xlsx
from .services.eligibility import EligibilityPolicy
from .services.pipeline.quality_gate import can_approve_lead, has_verified_contact_source
from .services.pipeline.source_audit import is_catalog_or_platform_source
from .services.pipeline.run_research import keyword_suggestions, run_research_pipeline
from .services.pipeline.scheduled_executor import (
    calculate_next_run_at,
    claim_scheduled_lead_task,
    execute_scheduled_lead_task,
)
from .services.pipeline.process_book import process_book
from .services.pipeline.manual_import_checks import build_imported_sales_brief, check_imported_lead_videos
from .services.social import harvest_social_profiles
from .services.verification import verify_lead_contacts
from .utils.normalize import normalized_author_key, normalized_book_key, normalize_text


logger = logging.getLogger(__name__)



HEADER_ALIASES = {
    "title": ["title", "book_title", "name", "book name", "book project", "book / project"],
    "author_name": ["author", "author_name", "author owner", "author / owner", "author / owner name", "author owner name"],
    "illustrator_name": ["illustrator", "illustrator_name"],
    "amazon_book_url": [
        "amazon_url",
        "amazon_book_url",
        "link",
        "amazon product url",
        "amazon book url",
        "amazon / book url",
        "amazon url",
        "amazon book / author url",
        "amazon book author url",
        "amazon / author url",
    ],
    "asin": ["asin"],
    "category": ["category"],
    "review_count": ["review_count", "no_of_ratings"],
    "rating": ["rating", "rating_out_of_5"],
    "publisher": ["publisher", "publisher imprint"],
    "publication_date": ["publication_date", "publication date"],
    "cover_image_url": ["cover_image_url"],
    "author_website": ["author_website", "author website", "official website", "website"],
    "contact_page_url": ["contact_page_url", "contact page url", "contact url", "contact form", "proof url", "email proof url", "email proof"],
    "public_email": ["public_email", "primary_email", "primary email", "email", "email address"],
    "representation_email": ["representation_email", "secondary_email", "secondary email", "agent email"],
    "publicist_email": ["publicist_email", "publicist email"],
    "public_phone": ["public_phone", "phone", "phone number", "phone dial format"],
    "public_email_source_url": ["public_email_source_url", "email source url", "email_source_url", "email proof url", "email proof"],
    "public_phone_source_url": ["public_phone_source_url", "phone source url", "phone_source_url", "phone proof url", "phone proof"],
    "verification_status": ["verification_status", "verification level", "lead quality status", "lead quality", "quality status"],
    "verification_score": ["verification_score"],
    "video_status": ["video_status"],
    "video_confidence": ["video_confidence"],
    "lead_score": ["lead_score"],
    "do_not_contact": ["do_not_contact", "dnc"],
    "location": ["location", "region", "phone country / region", "phone country region", "country"],
    "agent_name": ["agent_name", "agent name"],
    "publisher_url": ["publisher_url", "publisher url"],
    "instagram_url": ["instagram_url", "instagram url"],
    "facebook_url": ["facebook_url", "facebook url"],
    "tiktok_url": ["tiktok_url", "tiktok url"],
    "youtube_url": ["youtube_url", "youtube url"],
    "linkedin_url": ["linkedin_url", "linkedin url"],
    "sales_agent_summary": ["sales_agent_summary", "why this lead fits", "owner contact rationale", "contact rationale"],
    "suggested_first_line": ["suggested_first_line", "recommended first touch"],
    "suggested_pitch_angle": ["suggested_pitch_angle", "best fit bsp service"],
    "notes": ["notes"],
}


CONTACTABLE_Q = (
    Q(public_email__gt="")
    | Q(public_phone__gt="")
    | Q(representation_email__gt="")
    | Q(publicist_email__gt="")
    | Q(author_profile__contact_page_url__gt="")
)


def _is_safe_redirect_url(request, url: str) -> bool:
    """Only allow same-host redirects; block open-redirect abuse of next/REFERER."""
    return bool(url) and url_has_allowed_host_and_scheme(
        url,
        allowed_hosts={request.get_host()},
        require_https=request.is_secure(),
    )


def _contactable_queryset(queryset):
    return queryset.filter(CONTACTABLE_Q)


def _start_research_pipeline(research_run_id) -> None:
    def target() -> None:
        close_old_connections()
        try:
            run_research_pipeline(research_run_id)
        except Exception:
            logger.exception("Research pipeline thread crashed for run %s", research_run_id)
            try:
                run = ResearchRun.objects.get(pk=research_run_id)
                if run.status in {"pending", "running"}:
                    run.mark_failed("Background runner crashed; see server logs for details.")
            except Exception:
                logger.exception("Could not mark run %s as failed", research_run_id)
        finally:
            close_old_connections()

    threading.Thread(target=target, name=f"research-run-{research_run_id}", daemon=True).start()


def _start_manual_import_checks(research_run_id) -> None:
    """Run only the optional checks selected during a manual file import."""
    def target() -> None:
        close_old_connections()
        run = None
        try:
            run = ResearchRun.objects.get(pk=research_run_id)
            settings_json = run.settings_json or {}
            verify_contacts = bool(settings_json.get("verify_imported_contacts"))
            check_videos = bool(settings_json.get("run_video_search"))
            build_brief = bool(settings_json.get("run_groq_ai_extraction"))
            run.mark_running()
            batch = None
            if verify_contacts:
                batch = VerificationBatch.objects.create(
                    research_run=run,
                    status="running",
                    total_count=run.books.filter(leads__isnull=False).count(),
                    started_at=timezone.now(),
                    settings_json={"source": "manual_import", "check_network": True},
                )
            for lead in Lead.objects.filter(book__research_run=run).select_related("book", "author_profile").prefetch_related("evidence"):
                if verify_contacts:
                    verify_lead_contacts(lead, batch=batch, check_network=True)
                    batch.processed_count += 1
                    if lead.verification_status == "verified":
                        batch.verified_count += 1
                    elif lead.verification_status == "not_verified":
                        batch.not_verified_count += 1
                    else:
                        batch.other_count += 1
                    batch.save(update_fields=["processed_count", "verified_count", "not_verified_count", "other_count", "updated_at"])
                if check_videos:
                    check_imported_lead_videos(lead, use_ai=build_brief)
                if build_brief:
                    build_imported_sales_brief(lead, use_ai=True)
            if batch:
                batch.status = "completed"
                batch.completed_at = timezone.now()
                batch.save(update_fields=["status", "completed_at", "updated_at"])
            run.mark_completed()
        except Exception as exc:
            logger.exception("Manual import checks thread crashed for run %s", research_run_id)
            if run:
                run.mark_failed(str(exc))
        finally:
            close_old_connections()

    threading.Thread(target=target, name=f"manual-import-checks-{research_run_id}", daemon=True).start()


def _value(row: dict, field: str) -> str:
    for alias in HEADER_ALIASES[field]:
        key = _normalise_header(alias)
        if key in row and row[key]:
            return str(row[key]).strip()
    return ""


def _normalise_header(value) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().lower())).strip()


def _cell_text(value) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value).strip()


def _normalise_row(raw: dict) -> dict[str, str]:
    return {_normalise_header(key): _cell_text(value) for key, value in raw.items() if _normalise_header(key)}


def _find_header_row(values: list[tuple]) -> int | None:
    title_aliases = {_normalise_header(alias) for alias in HEADER_ALIASES["title"]}
    author_aliases = {_normalise_header(alias) for alias in HEADER_ALIASES["author_name"]}
    known_aliases = {
        _normalise_header(alias)
        for aliases in HEADER_ALIASES.values()
        for alias in aliases
    }
    best: tuple[int, int] | None = None
    for index, row in enumerate(values):
        headers = {_normalise_header(value) for value in row if _normalise_header(value)}
        if not (headers & title_aliases and headers & author_aliases):
            continue
        score = len(headers & known_aliases)
        if best is None or score > best[1]:
            best = (index, score)
    return best[0] if best else None


def parse_import_file(uploaded_file) -> tuple[list[dict[str, str]], str]:
    """Read a CSV/XLSX/XLSM upload and locate a lead-table header automatically."""
    filename = uploaded_file.name or "uploaded file"
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    payload = uploaded_file.read()
    if suffix == "csv":
        try:
            content = payload.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("CSV files must be UTF-8 encoded.") from exc
        rows = [_normalise_row(row) for row in csv.DictReader(io.StringIO(content))]
        if not rows:
            raise ValueError("The file does not contain any data rows.")
        return rows, "CSV"

    try:
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("The Excel file could not be read. Upload a valid XLSX or XLSM workbook.") from exc

    parsed_sheets: dict[str, list[dict[str, str]]] = {}
    for worksheet in workbook.worksheets:
        sampled = list(worksheet.iter_rows(min_row=1, max_row=min(25, worksheet.max_row), values_only=True))
        header_index = _find_header_row(sampled)
        if header_index is not None:
            headers = [_normalise_header(value) for value in sampled[header_index]]
        else:
            headers = [_normalise_header(value) for value in sampled[0]] if sampled else []
        if not headers:
            continue
        rows = []
        start_row = (header_index + 2) if header_index is not None else 2
        for values in worksheet.iter_rows(min_row=start_row, values_only=True):
            raw = {headers[index]: _cell_text(value) for index, value in enumerate(values) if index < len(headers) and headers[index]}
            if any(raw.values()):
                rows.append(raw)
        parsed_sheets[worksheet.title] = rows
        if header_index is not None and rows:
            return rows, f"Excel: {worksheet.title}"

    # The app's formatted XLSX export stores authors, books, contact evidence,
    # and outreach fields on separate tabs. Reassemble those tabs by Lead ID so
    # an exported workbook can be imported again without manual reformatting.
    author_rows = next((rows for rows in parsed_sheets.values() if rows and {"lead id", "author name"}.issubset(rows[0])), [])
    book_rows = next((rows for rows in parsed_sheets.values() if rows and {"lead id", "book title"}.issubset(rows[0])), [])
    if author_rows and book_rows:
        authors = {row.get("lead id", ""): row for row in author_rows if row.get("lead id")}
        contacts = next((rows for rows in parsed_sheets.values() if rows and {"lead id", "contact type", "contact value", "source url"}.issubset(rows[0])), [])
        outreach = next((rows for rows in parsed_sheets.values() if rows and {"lead id", "outreach angle"}.issubset(rows[0])), [])
        contact_map: dict[str, list[dict[str, str]]] = {}
        for item in contacts:
            contact_map.setdefault(item.get("lead id", ""), []).append(item)
        outreach_map = {row.get("lead id", ""): row for row in outreach if row.get("lead id")}
        merged = []
        for book in book_rows:
            lead_id = book.get("lead id", "")
            row = {**book, **authors.get(lead_id, {}), **outreach_map.get(lead_id, {})}
            for contact in contact_map.get(lead_id, []):
                contact_type = normalize_text(contact.get("contact type", ""))
                if "phone" in contact_type:
                    row.setdefault("public phone", contact.get("contact value", ""))
                    row.setdefault("phone source url", contact.get("source url", ""))
                elif "email" in contact_type:
                    row.setdefault("public email", contact.get("contact value", ""))
                    row.setdefault("email source url", contact.get("source url", ""))
            merged.append(row)
        if merged:
            return merged, "Excel: formatted app export"
    raise ValueError("No lead table was found. Include title/book title and author columns within the first 25 rows of a sheet.")


def _int_or_none(value: str):
    """Parse integer fields tolerating locale thousands separators.

    "1,234" and "1.234" both mean 1234; a plain decimal like "4.5" is
    truncated to 4 rather than misread as 45 or 4500.
    """
    try:
        cleaned = str(value).strip()
        if not cleaned:
            return None
        if re.fullmatch(r"\d{1,3}([.,]\d{3})+", cleaned):
            return int(re.sub(r"[.,]", "", cleaned))
        return int(float(cleaned.replace(",", "")))
    except (ValueError, TypeError):
        return None


def _decimal_or_none(value: str):
    try:
        raw = str(value).strip()
        if not raw:
            return None
        dec = Decimal(raw).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        if dec < Decimal("0") or dec > Decimal("5.00"):
            return None
        return dec
    except (InvalidOperation, ValueError, TypeError):
        return None


def _float_or_zero(value: str) -> float:
    try:
        return float(str(value).strip() or 0)
    except (TypeError, ValueError):
        return 0.0



def _bool_value(value: str) -> bool:
    return normalize_text(value) in {"1", "true", "yes", "y", "on", "do not contact", "dnc"}


def _status_value(value: str) -> str:
    norm = normalize_text(value)
    if not norm:
        return "other"
    if norm in {"not verified", "invalid", "failed", "undeliverable", "no"}:
        return "not_verified"
    if (
        norm in {"verified", "valid", "approved", "complete", "yes", "a", "a+", "a-"}
        or "verified" in norm
        or norm.startswith(("a ", "a -", "a-", "a+", "b ", "b -", "b-", "b+", "high", "complete"))
    ) and not any(kw in norm for kw in ("needs proof", "pending", "not verified", "unverified", "incomplete", "no contact")):
        return "verified"
    return "other"


def _video_status_value(value: str) -> str:
    value = normalize_text(value).replace(" ", "_")
    allowed = {choice for choice, _ in Lead.VIDEO_STATUS_CHOICES}
    return value if value in allowed else "not_checked"


def _source_url(row: dict, preferred_field: str, fallback: str = "") -> str:
    value = _value(row, preferred_field) or fallback
    return value if value.startswith(("https://", "http://")) else ""


@transaction.atomic
def import_books_from_rows(rows: list[dict], run: ResearchRun, source_label: str = "uploaded_file/manual") -> int:
    """Import rows atomically so a mid-file failure never leaves a half-imported run."""
    count = 0
    is_manual_import = run.source_provider == "manual"
    for raw in rows:
        row = _normalise_row(raw)
        title = _value(row, "title")
        author_name = _value(row, "author_name")
        if not title:
            continue
        amazon_url = _value(row, "amazon_book_url")
        asin = (_value(row, "asin") or extract_asin(amazon_url) or "").upper()
        if asin and not amazon_url:
            amazon_url = normalize_amazon_book_url(f"https://www.amazon.com/dp/{asin}", os.getenv("AMAZON_ASSOCIATE_TAG") or None)
        elif amazon_url and asin:
            amazon_url = normalize_amazon_book_url(amazon_url, os.getenv("AMAZON_ASSOCIATE_TAG") or None)
        book = Book.objects.create(
            research_run=run,
            title=title,
            author_name=author_name,
            illustrator_name=_value(row, "illustrator_name"),
            asin=asin,
            amazon_book_url=amazon_url,
            amazon_source_url=amazon_url or source_label,
            category=_value(row, "category"),
            review_count=_int_or_none(_value(row, "review_count")),
            rating=_decimal_or_none(_value(row, "rating")),
            publisher=_value(row, "publisher"),
            publication_date=_value(row, "publication_date"),
            cover_image_url=_value(row, "cover_image_url"),
            normalized_key=normalized_book_key(title, author_name, asin),
            book_data_confidence=0.9,
            source_provider=run.source_provider,
            source_raw_json={"source": source_label},
        )
        if amazon_url:
            Evidence.objects.create(
                book=book,
                evidence_type="manual" if run.source_provider == "manual" else "amazon_search_result",
                field_name="amazon_book_url",
                field_value=amazon_url,
                source_url=source_label,
                confidence=0.9,
                is_primary=True,
            )
        website = _source_url(row, "author_website")
        contact_page_url = _source_url(row, "contact_page_url", website)
        # Reuse an existing profile for the same author instead of creating a
        # duplicate profile for every row; only fill still-blank fields.
        profile_defaults = {
            "author_name": author_name or "Unknown author",
            "canonical_website": website,
            "contact_page_url": contact_page_url,
            "publisher_url": _source_url(row, "publisher_url"),
            "instagram_url": _source_url(row, "instagram_url"),
            "facebook_url": _source_url(row, "facebook_url"),
            "tiktok_url": _source_url(row, "tiktok_url"),
            "youtube_url": _source_url(row, "youtube_url"),
            "linkedin_url": _source_url(row, "linkedin_url"),
            "location": _value(row, "location"),
            "agent_name": _value(row, "agent_name"),
            "representation_email": _value(row, "representation_email"),
            "publicist_email": _value(row, "publicist_email"),
            "identity_confidence": 0.7 if author_name else 0.1,
            "identity_reason": "Manually imported from a client-supplied lead file; source-backed contact verification is optional.",
        }
        author_key = normalized_author_key(author_name)
        try:
            author_profile, profile_created = AuthorProfile.objects.get_or_create(
                normalized_author_key=author_key,
                defaults=profile_defaults,
            )
        except AuthorProfile.MultipleObjectsReturned:
            # Legacy databases can already hold duplicate keys; reuse the oldest.
            author_profile = AuthorProfile.objects.filter(normalized_author_key=author_key).order_by("created_at").first()
            profile_created = False
        if not profile_created:
            profile_updates = {}
            for field, new_value in {
                "canonical_website": website,
                "contact_page_url": contact_page_url,
                "publisher_url": _source_url(row, "publisher_url"),
                "instagram_url": _source_url(row, "instagram_url"),
                "facebook_url": _source_url(row, "facebook_url"),
                "tiktok_url": _source_url(row, "tiktok_url"),
                "youtube_url": _source_url(row, "youtube_url"),
                "linkedin_url": _source_url(row, "linkedin_url"),
                "location": _value(row, "location"),
                "agent_name": _value(row, "agent_name"),
                "representation_email": _value(row, "representation_email"),
                "publicist_email": _value(row, "publicist_email"),
            }.items():
                if new_value and not getattr(author_profile, field):
                    profile_updates[field] = new_value
            if profile_updates:
                for field, new_value in profile_updates.items():
                    setattr(author_profile, field, new_value)
                author_profile.save(update_fields=[*profile_updates, "updated_at"])
        imported_verification_score = min(100, _int_or_none(_value(row, "verification_score")) or 0)
        if is_manual_import:
            # An upload records who attested to a value; it is never evidence
            # that the value passed source, identity, and deliverability checks.
            imported_verification_status = "other"
            imported_verification_score = 0
            imported_verification_reason = "Imported contact values are awaiting system verification."
            imported_verification_reasons = ["uploader_attested", "import_awaiting_verification"]
            imported_verified_at = None
        else:
            imported_verification_status = _status_value(_value(row, "verification_status"))
            imported_verification_reason = ""
            imported_verification_reasons = []
            imported_verified_at = None
        lead = Lead.objects.create(
            book=book,
            author_profile=author_profile,
            public_email=_value(row, "public_email"),
            public_phone=_value(row, "public_phone"),
            representation_email=_value(row, "representation_email"),
            publicist_email=_value(row, "publicist_email"),
            location=_value(row, "location"),
            lead_score=_int_or_none(_value(row, "lead_score")) or 0,
            verification_score=imported_verification_score,
            verification_status=imported_verification_status,
            verification_reason=imported_verification_reason,
            verification_reasons_json=imported_verification_reasons,
            verification_version="manual-import-v2" if is_manual_import else "",
            verified_at=imported_verified_at,
            uploader_attested=is_manual_import,
            uploader_attested_at=timezone.now() if is_manual_import else None,
            video_status=_video_status_value(_value(row, "video_status")),
            video_confidence=min(1, max(0, _float_or_zero(_value(row, "video_confidence")))),
            sales_agent_summary=_value(row, "sales_agent_summary"),
            suggested_pitch_angle=_value(row, "suggested_pitch_angle"),
            suggested_first_line=_value(row, "suggested_first_line"),
            do_not_contact=_bool_value(_value(row, "do_not_contact")),
            notes=_value(row, "notes"),
            manual_review_status="needs_review",
        )
        contact_sources = {
            "public_email": _source_url(row, "public_email_source_url", contact_page_url or website),
            "public_phone": _source_url(row, "public_phone_source_url", contact_page_url or website),
            "representation_email": contact_page_url or website,
            "publicist_email": contact_page_url or website,
        }
        primary_contact = None
        contact_roles = {
            "public_email": ("email", "author"),
            "representation_email": ("email", "agent"),
            "publicist_email": ("email", "publicist"),
            "public_phone": ("phone", "author"),
        }
        for field_name, source_url in contact_sources.items():
            value = getattr(lead, field_name)
            if value and is_manual_import:
                channel, role = contact_roles[field_name]
                candidate = ContactCandidate.objects.create(
                    lead=lead,
                    channel=channel,
                    role=role,
                    raw_value=value,
                    normalized_value=value,
                    verification_status="other",
                    verification_score=0,
                    deliverability_status="not_checked",
                    is_primary=primary_contact is None,
                    selected_reason="Uploader-attested import; system verification has not run.",
                )
                if primary_contact is None:
                    primary_contact = candidate
            if value and source_url:
                Evidence.objects.create(
                    lead=lead,
                    book=book,
                    author_profile=author_profile,
                    evidence_type="manual",
                    field_name=field_name,
                    field_value=value,
                    source_url=source_url,
                    source_title="Manual import",
                    confidence=0.5 if is_manual_import else 0.7,
                    is_primary=True,
                )
        if primary_contact:
            lead.primary_contact = primary_contact
            lead.save(update_fields=["primary_contact", "updated_at"])
        count += 1
    return count


def dashboard(request):
    leads = lead_queryset_for_user(request.user)
    failed_runs = ResearchRun.objects.filter(status="failed")
    total_leads = leads.count()
    verified_leads = EligibilityPolicy.verified_ready(leads).count()
    not_verified_leads = leads.filter(verification_status="not_verified").count()
    other_leads = EligibilityPolicy.needs_review(leads).count()
    public_emails = leads.exclude(public_email="").count()

    def percentage(value):
        return round((value / total_leads) * 100) if total_leads else 0

    assignments = LeadAssignment.objects.filter(is_current=True)
    if user_role(request.user) == "sales":
        assignments = assignments.filter(assigned_to=request.user)
    context = {
        "total_leads": total_leads,
        "verified_leads": verified_leads,
        "not_verified_leads": not_verified_leads,
        "other_leads": other_leads,
        "public_emails": public_emails,
        "verified_percent": percentage(verified_leads),
        "not_verified_percent": percentage(not_verified_leads),
        "other_percent": percentage(other_leads),
        "contact_percent": percentage(public_emails),
        "amazon_urls": Book.objects.exclude(amazon_book_url="").count(),
        "no_public_video_found": leads.filter(video_status="no_public_video_found").count(),
        "video_found": leads.filter(video_status__in=["found_trailer", "found_animated_video", "found_read_aloud_only"]).count(),
        "needs_review": leads.filter(manual_review_status="needs_review").count(),
        "approved": leads.filter(manual_review_status="approved").count(),
        "do_not_contact": leads.filter(do_not_contact=True).count(),
        "failed_runs": failed_runs.count(),
        "latest_errors": failed_runs.exclude(error_message="").order_by("-completed_at", "-created_at")[:3],
        "latest_runs": ResearchRun.objects.order_by("-created_at")[:8],
        "latest_verification_batches": VerificationBatch.objects.order_by("-created_at")[:5],
        "assignment_stats": {
            "total": assignments.count(),
            "pending": assignments.filter(status="pending").count(),
            "in_progress": assignments.filter(status__in=["in_progress", "contacted", "follow_up"]).count(),
            "completed": assignments.filter(status="completed").count(),
            "converted": assignments.filter(status="converted").count(),
            "invalid": assignments.filter(status__in=["invalid_data", "duplicate"]).count(),
        },
        "team_performance": LeadAssignment.objects.filter(is_current=True)
            .values("assigned_to__id", "assigned_to__first_name", "assigned_to__last_name", "assigned_to__username")
            .annotate(
                total=Count("id"),
                completed=Count("id", filter=Q(status="completed")),
                converted=Count("id", filter=Q(status="converted")),
                pending=Count("id", filter=Q(status="pending")),
            ).order_by("-converted", "-completed") if is_management(request.user) else [],
    }
    return render(request, "leadfinder/dashboard.html", context)


@roles_required("super_admin", "admin")
@require_POST
def delete_rejected_leads(request):
    deleted_count, _ = Lead.objects.filter(
        Q(lead_tier="rejected") | Q(manual_review_status="rejected")
    ).delete()
    messages.success(request, f"Successfully deleted {deleted_count} rejected leads.")
    return redirect("leadfinder:dashboard")


def _start_scheduled_task(task_id) -> bool:
    if not claim_scheduled_lead_task(task_id, force=True):
        return False

    def target() -> None:
        close_old_connections()
        try:
            execute_scheduled_lead_task(task_id, force=True, preclaimed=True)
        except Exception:
            logger.exception("Scheduled lead task thread crashed for task %s", task_id)
            try:
                ScheduledLeadTask.objects.filter(pk=task_id, execution_status="running").update(
                    execution_status="failed",
                    last_error="Background runner crashed; see server logs for details.",
                )
            except Exception:
                logger.exception("Could not mark scheduled task %s as failed", task_id)
        finally:
            close_old_connections()

    threading.Thread(target=target, name=f"scheduled-lead-task-{task_id}", daemon=True).start()
    return True


@roles_required("super_admin", "admin")
def scheduled_task_list(request):
    latest_runs = ResearchRun.objects.filter(scheduled_task=OuterRef("pk")).order_by("-created_at")
    tasks = ScheduledLeadTask.objects.annotate(
        latest_run_id=Subquery(latest_runs.values("id")[:1]),
        latest_run_status=Subquery(latest_runs.values("status")[:1]),
    )
    now = timezone.now()
    return render(
        request,
        "leadfinder/scheduled_task_list.html",
        {
            "tasks": tasks,
            "stats": {
                "total": tasks.count(),
                "active": tasks.filter(is_active=True).count(),
                "running": tasks.filter(execution_status="running").count(),
                "due": tasks.filter(is_active=True, next_run_at__lte=now).exclude(execution_status="running").count(),
            },
            "now": now,
        },
    )


@roles_required("super_admin", "admin")
def scheduled_task_create(request):
    if request.method == "POST":
        form = ScheduledLeadTaskForm(request.POST)
        if form.is_valid():
            task = form.save(commit=False)
            task.next_run_at = calculate_next_run_at(task)
            task.save()
            messages.success(request, f'“{task.name}” is scheduled and ready.')
            return redirect("leadfinder:scheduled_task_list")
    else:
        form = ScheduledLeadTaskForm()
    return render(
        request,
        "leadfinder/scheduled_task_form.html",
        {"form": form, "page_title": "Create scheduled hunt", "submit_label": "Create schedule"},
    )


@roles_required("super_admin", "admin")
def scheduled_task_edit(request, pk):
    task = get_object_or_404(ScheduledLeadTask, pk=pk)
    if request.method == "POST":
        form = ScheduledLeadTaskForm(request.POST, instance=task)
        if form.is_valid():
            task = form.save(commit=False)
            task.next_run_at = calculate_next_run_at(task)
            task.save()
            messages.success(request, f'“{task.name}” was updated.')
            return redirect("leadfinder:scheduled_task_list")
    else:
        form = ScheduledLeadTaskForm(instance=task)
    return render(
        request,
        "leadfinder/scheduled_task_form.html",
        {
            "form": form,
            "task": task,
            "page_title": "Edit scheduled hunt",
            "submit_label": "Save changes",
        },
    )


@roles_required("super_admin", "admin")
@require_POST
def scheduled_task_toggle(request, pk):
    task = get_object_or_404(ScheduledLeadTask, pk=pk)
    task.is_active = not task.is_active
    update_fields = ["is_active", "updated_at"]
    if task.is_active:
        task.next_run_at = calculate_next_run_at(task)
        update_fields.append("next_run_at")
        messages.success(request, f'“{task.name}” resumed. Next run: {task.next_run_at:%b %d, %Y %H:%M}.')
    else:
        messages.info(request, f'“{task.name}” is paused. A currently running job will finish safely.')
    task.save(update_fields=update_fields)
    return redirect("leadfinder:scheduled_task_list")


@roles_required("super_admin", "admin")
@require_POST
def scheduled_task_run_now(request, pk):
    task = get_object_or_404(ScheduledLeadTask, pk=pk)
    if task.execution_status == "running":
        messages.warning(request, f'“{task.name}” is already running.')
    else:
        started = _start_scheduled_task(task.id)
        if started is False:
            messages.warning(request, f'“{task.name}” is already running.')
        else:
            messages.success(request, f'“{task.name}” is starting in the background.')
    return redirect("leadfinder:scheduled_task_list")


@roles_required("super_admin", "admin")
@require_POST
def scheduled_task_delete(request, pk):
    task = get_object_or_404(ScheduledLeadTask, pk=pk)
    if task.execution_status == "running":
        messages.warning(request, "Pause future runs after the current execution finishes, then delete the schedule.")
        return redirect("leadfinder:scheduled_task_list")
    name = task.name
    task.delete()
    messages.success(request, f'“{name}” was deleted. Existing research runs and leads were preserved.')
    return redirect("leadfinder:scheduled_task_list")



@roles_required("super_admin", "admin")
def run_new(request):
    if request.method == "POST":
        form = ResearchRunForm(request.POST)
        if form.is_valid():
            if form.cleaned_data["source_provider"] == "csv":
                messages.info(request, "CSV runs start from the import screen.")
                return redirect("leadfinder:import_csv")
            run = ResearchRun.objects.create(
                keyword=form.cleaned_data["keyword"],
                source_provider=form.cleaned_data["source_provider"],
                marketplace=form.cleaned_data["marketplace"],
                max_books=form.cleaned_data["max_books"],
                settings_json={
                    "require_amazon_url": form.cleaned_data["require_amazon_url"],
                    "require_public_email": form.cleaned_data["require_public_email"],
                    "include_social_only_leads": form.cleaned_data["include_social_only_leads"],
                    "run_video_search": form.cleaned_data["run_video_search"],
                    "run_groq_ai_extraction": form.cleaned_data["run_groq_ai_extraction"],
                    "location_preference": form.cleaned_data["location_preference"],
                },
            )
            if run.source_provider != "csv":
                _start_research_pipeline(run.id)
            messages.success(request, "Research run created. Search is running in the background.")
            return redirect("leadfinder:run_detail", pk=run.id)
    else:
        form = ResearchRunForm()
    return render(request, "leadfinder/run_new.html", {"form": form, "keyword_suggestions": keyword_suggestions()})


@roles_required("super_admin", "admin")
def booklife_run(request):
    category_labels = {category.slug or "all": category.label for category in BOOKLIFE_CATEGORIES}
    if request.method == "POST":
        form = BookLifeRunForm(request.POST)
        if form.is_valid():
            categories = form.cleaned_data["booklife_categories"]
            labels = [category_labels.get(slug, slug) for slug in categories]
            keyword = "BookLife: " + ", ".join(labels[:3])
            if len(labels) > 3:
                keyword = f"{keyword} + {len(labels) - 3} more"
            run = ResearchRun.objects.create(
                keyword=keyword,
                source_provider="booklife",
                marketplace="US",
                max_books=form.cleaned_data["max_books"],
                settings_json={
                    "booklife_categories": categories,
                    "booklife_category_labels": labels,
                    "booklife_age_filter": form.cleaned_data["age_filter"],
                    "enrichment_provider": form.cleaned_data["enrichment_provider"],
                    "require_amazon_url": False,
                    "require_public_email": form.cleaned_data["require_public_email"],
                    "include_social_only_leads": form.cleaned_data["include_social_only_leads"],
                    "run_video_search": form.cleaned_data["run_video_search"],
                    "run_groq_ai_extraction": form.cleaned_data["run_groq_ai_extraction"],
                },
            )
            _start_research_pipeline(run.id)
            messages.success(request, "BookLife runner created. Discovery and enrichment are running in the background.")
            return redirect("leadfinder:run_detail", pk=run.id)
    else:
        form = BookLifeRunForm(initial={"booklife_categories": ["all"]})
    return render(
        request,
        "leadfinder/booklife_run.html",
        {
            "form": form,
            "category_groups": grouped_booklife_categories(),
            "selected_categories": form["booklife_categories"].value() or ["all"],
        },
    )


@roles_required("super_admin", "admin")
def run_list(request):
    base_runs = ResearchRun.objects.annotate(
        total_books=Count("books", distinct=True),
        total_leads=Count("books__leads", distinct=True),
    )
    total_count = base_runs.count()
    completed_count = base_runs.filter(status="completed").count()
    running_count = base_runs.filter(status__in=["pending", "running"]).count()
    failed_count = base_runs.filter(status__in=["failed", "canceled"]).count()

    runs = base_runs
    query = request.GET.get("q", "").strip()
    provider = request.GET.get("provider", "").strip()
    status = request.GET.get("status", "").strip()
    if query:
        runs = runs.filter(Q(keyword__icontains=query) | Q(source_provider__icontains=query))
    if provider:
        runs = runs.filter(source_provider=provider)
    if status == "active":
        runs = runs.filter(status__in=["pending", "running"])
    elif status == "failed":
        runs = runs.filter(status__in=["failed", "canceled"])
    elif status in {"completed", "pending", "running", "canceled"}:
        runs = runs.filter(status=status)
    runs = runs.order_by("-created_at")
    
    # Pagination
    try:
        per_page = int(request.GET.get("per_page", 15))
    except ValueError:
        per_page = 15
    if per_page not in [15, 25, 50, 100]:
        per_page = 15

    paginator = Paginator(runs, per_page)
    page_number = request.GET.get("page", 1)
    page_obj = paginator.get_page(page_number)

    return render(
        request, 
        "leadfinder/run_list.html", 
        {
            "runs": runs,
            "page_obj": page_obj,
            "per_page": per_page,
            "stats": {
                "total": total_count,
                "completed": completed_count,
                "running": running_count,
                "failed": failed_count,
            },
            "active_run_filter": status,
        }
    )


@roles_required("super_admin", "admin")
def run_detail(request, pk):
    run = get_object_or_404(ResearchRun, pk=pk)
    books = list(run.books.prefetch_related("leads", "evidence").all())
    leads = list(
        Lead.objects.filter(book__research_run=run)
        .select_related("book", "author_profile", "brief")
        .prefetch_related("evidence", "contact_candidates")
    )
    search_logs = list(run.search_logs.annotate(
        total_results=Count("results"),
        amazon_results=Count("results", filter=Q(results__classification="amazon_book")),
    ).order_by("created_at"))
    no_search_results = bool(search_logs) and all(log.total_results == 0 for log in search_logs)
    settings_json = run.settings_json or {}
    groq_evidence_count = Evidence.objects.filter(
        lead__book__research_run=run,
        evidence_type="groq_extraction",
    ).count()
    ai_brief_count = sum(1 for lead in leads if hasattr(lead, "brief"))
    video_evidence_count = Lead.objects.filter(book__research_run=run, videos__isnull=False).distinct().count()
    books_missing_author_count = sum(1 for book in books if not (book.author_name or "").strip())
    leads_missing_contact_count = sum(
        1
        for lead in leads
        if not any(
            [
                lead.public_email,
                lead.public_phone,
                lead.representation_email,
                lead.publicist_email,
                lead.primary_contact_id,
            ]
        )
    )
    ai_diagnostics = {
        "enabled": bool(settings_json.get("run_groq_ai_extraction", True)),
        "video_enabled": bool(settings_json.get("run_video_search", True)),
        "brief_count": ai_brief_count,
        "groq_evidence_count": groq_evidence_count,
        "video_evidence_count": video_evidence_count,
        "books_missing_author_count": books_missing_author_count,
        "leads_missing_contact_count": leads_missing_contact_count,
        "total_books": len(books),
        "total_leads": len(leads),
    }
    lead_checks = []
    for lead in leads:
        passed, errors = can_approve_lead(
            lead,
            allow_incomplete_video=not ai_diagnostics["video_enabled"],
        )
        lead_checks.append({"lead": lead, "passed": passed, "errors": errors})
    book_errors = [
        {
            "book": book,
            "errors": (book.source_raw_json or {}).get("book_processing_errors", []),
        }
        for book in books
        if (book.source_raw_json or {}).get("book_processing_errors")
    ]
    return render(
        request,
        "leadfinder/run_detail.html",
        {
            "run": run,
            "books": books,
            "leads": leads,
            "search_logs": search_logs,
            "lead_checks": lead_checks,
            "book_errors": book_errors,
            "no_search_results": no_search_results,
            "ai_diagnostics": ai_diagnostics,
        },
    )


@roles_required("super_admin", "admin")
def run_agent_status(request, pk):
    run = get_object_or_404(ResearchRun, pk=pk)
    books = run.books.all()
    total_books = books.count()
    completed_books = 0
    for book in books:
        status = (book.source_raw_json or {}).get("processing_status", "")
        if status in {"completed", "failed"}:
            completed_books += 1
            
    progress = 0
    if run.status == "completed":
        progress = 100
    elif run.status == "failed":
        progress = 100
    elif total_books > 0:
        progress = int((completed_books / total_books) * 100)
        if progress >= 100 and run.status == "running":
            progress = 99
            
    discovered_count = total_books
    verified_leads_count = EligibilityPolicy.verified_ready(
        Lead.objects.filter(book__research_run=run)
    ).count()
    
    thoughts = run.agent_thoughts.order_by("created_at")
    recent_thoughts = [
        {
            "agent_name": t.agent_name,
            "message": t.message,
            "created_at": t.created_at.strftime("%I:%M:%S %p"),
        }
        for t in thoughts
    ]
    
    active_agent = "Scout"
    if run.status in {"completed", "failed", "canceled"}:
        active_agent = "None"
    elif recent_thoughts:
        active_agent = recent_thoughts[-1]["agent_name"]
        
    return JsonResponse({
        "status": run.get_status_display() if hasattr(run, "get_status_display") else run.status,
        "run_status_raw": run.status,
        "progress": progress,
        "discovered_count": discovered_count,
        "verified_leads_count": verified_leads_count,
        "active_agent": active_agent,
        "thoughts": recent_thoughts,
    })


@roles_required("super_admin", "admin")
def import_csv(request):
    preview = []
    import_summary = None
    if request.method == "POST":
        form = CSVImportForm(request.POST, request.FILES)
        if form.is_valid():
            upload = form.cleaned_data["csv_file"]
            try:
                rows, source_description = parse_import_file(upload)
            except ValueError as exc:
                form.add_error("csv_file", str(exc))
            else:
                preview = rows[:20]
                import_summary = {"row_count": len(rows), "source_description": source_description, "headers": list(rows[0]) if rows else []}
                if "preview" not in request.POST:
                    checks_requested = any(
                        (
                            form.cleaned_data["run_video_search"],
                            form.cleaned_data["verify_imported_contacts"],
                            form.cleaned_data["run_groq_ai_extraction"],
                        )
                    )
                    run = ResearchRun.objects.create(
                        keyword=f"Manual import: {upload.name}",
                        source_provider="manual",
                        max_books=max(len(rows), 1),
                        status="pending" if checks_requested else "completed",
                        completed_at=None if checks_requested else timezone.now(),
                        settings_json={
                            "manual_import": True,
                            "source_file_name": upload.name,
                            "source_description": source_description,
                            "require_public_email": False,
                            "run_video_search": form.cleaned_data["run_video_search"],
                            "verify_imported_contacts": form.cleaned_data["verify_imported_contacts"],
                            "run_groq_ai_extraction": form.cleaned_data["run_groq_ai_extraction"],
                        },
                    )
                    imported = import_books_from_rows(rows, run, source_label=upload.name)
                    run.max_books = imported
                    run.save(update_fields=["max_books"])
                    if checks_requested:
                        _start_manual_import_checks(run.id)
                        messages.success(request, f"Imported {imported} manual leads. Only the selected checks are running in the background.")
                    else:
                        messages.success(request, f"Imported {imported} manual leads. No AI or external checks were run.")
                    return redirect("leadfinder:run_detail", pk=run.id)
    else:
        form = CSVImportForm()
    return render(request, "leadfinder/import_csv.html", {"form": form, "preview": preview, "import_summary": import_summary})


def lead_list(request):
    form = LeadFilterForm(request.GET or None)
    verified_contacts = ContactCandidate.objects.filter(
        lead_id=OuterRef("pk"),
        verification_status="verified",
    )
    current_assignment_prefetch = Prefetch(
        "assignments",
        queryset=LeadAssignment.objects.filter(is_current=True).select_related("assigned_to"),
        to_attr="current_assignments",
    )
    leads = EligibilityPolicy.annotate(lead_queryset_for_user(
        request.user,
        Lead.objects.select_related("book", "author_profile", "primary_contact").prefetch_related(current_assignment_prefetch),
    )).annotate(
        has_verified_email=Exists(verified_contacts.filter(channel="email")),
        has_verified_phone=Exists(verified_contacts.filter(channel="phone")),
    )

    email_q = Q(public_email__gt="") | Q(representation_email__gt="") | Q(publicist_email__gt="")
    phone_q = Q(public_phone__gt="")

    # Apply standard filters
    if form.is_valid():
        cd = form.cleaned_data

        # Keyword Search
        if cd.get("q"):
            q_term = cd["q"].strip()
            leads = leads.filter(
                Q(book__title__icontains=q_term) |
                Q(book__author_name__icontains=q_term) |
                Q(public_email__icontains=q_term) |
                Q(representation_email__icontains=q_term) |
                Q(publicist_email__icontains=q_term)
            )
            
        if cd.get("valid_only"):
            leads = EligibilityPolicy.verified_ready(leads)
        if cd.get("verification_status"):
            if cd["verification_status"] == "verified":
                leads = EligibilityPolicy.verified_ready(leads)
            elif cd["verification_status"] == "not_verified":
                leads = leads.filter(verification_status="not_verified")
            elif cd["verification_status"] == "other":
                leads = EligibilityPolicy.needs_review(leads)
        if cd.get("video_status"):
            leads = leads.filter(video_status=cd["video_status"])
        if cd.get("score_min") is not None:
            leads = leads.filter(verification_score__gte=cd["score_min"])
        if cd.get("score_max") is not None:
            leads = leads.filter(verification_score__lte=cd["score_max"])
        if cd.get("has_email"):
            leads = leads.filter(email_q | Q(has_verified_email=True))
        if cd.get("verified_email"):
            leads = leads.filter(has_verified_email=True)
        if cd.get("has_phone"):
            leads = leads.filter(phone_q | Q(has_verified_phone=True))
        if cd.get("verified_phone"):
            leads = leads.filter(has_verified_phone=True)
        if cd.get("has_amazon_url"):
            leads = leads.exclude(book__amazon_book_url="")
        if cd.get("has_website"):
            leads = leads.exclude(author_profile__canonical_website="")
        if cd.get("do_not_contact"):
            leads = leads.filter(do_not_contact=True)
        if cd.get("min_confidence"):
            leads = leads.filter(extraction_confidence__gte=float(cd["min_confidence"]))
        if cd.get("pub_year_start") is not None:
            leads = leads.exclude(book__publication_date="").filter(book__publication_date__gte=str(cd["pub_year_start"]))
        if cd.get("pub_year_end") is not None:
            leads = leads.exclude(book__publication_date="").filter(book__publication_date__lt=str(int(cd["pub_year_end"]) + 1))
        if cd.get("assignment_status"):
            leads = leads.filter(assignments__is_current=True, assignments__status=cd["assignment_status"])
        if cd.get("assigned_to") and is_management(request.user):
            leads = leads.filter(assignments__is_current=True, assignments__assigned_to=cd["assigned_to"])

    # Handle sorting
    order_field = "-created_at"  # Default to Newest First as requested by user
    if form.is_valid() and cd.get("sort_by"):
        order_field = cd["sort_by"]
    leads = leads.order_by(order_field)
    
    # Calculate stats for dashboard header
    stats_queryset = EligibilityPolicy.annotate(lead_queryset_for_user(request.user)).annotate(
        has_verified_email=Exists(verified_contacts.filter(channel="email")),
        has_verified_phone=Exists(verified_contacts.filter(channel="phone")),
    )
    stats_total = stats_queryset.count()
    stats_verified = EligibilityPolicy.verified_ready(stats_queryset).count()
    stats_verified_email = stats_queryset.filter(has_verified_email=True).count()
    stats_verified_phone = stats_queryset.filter(has_verified_phone=True).count()
    stats_other = EligibilityPolicy.needs_review(stats_queryset).count()
    stats_email = stats_verified_email
    
    # Pagination
    try:
        per_page = int(request.GET.get('per_page', 25))
    except ValueError:
        per_page = 25
    if per_page not in [25, 50, 100]:
        per_page = 25
        
    paginator = Paginator(leads, per_page)
    page_number = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_number)
    
    context = {
        "form": form,
        "page_obj": page_obj,
        "per_page": per_page,
        "leads": leads,
        "stats": {
            "total": stats_total,
            "verified": stats_verified,
            "other": stats_other,
            "email": stats_email,
            "verified_email": stats_verified_email,
            "verified_phone": stats_verified_phone,
        },
        "salespeople": get_user_model().objects.filter(
            is_active=True, leadfinder_profile__role="sales", leadfinder_profile__is_available_for_assignment=True
        ).order_by("first_name", "last_name", "username") if is_management(request.user) else [],
    }
    return render(request, "leadfinder/lead_list.html", context)



def lead_detail(request, pk):
    from collections import Counter

    lead = get_object_or_404(
        lead_queryset_for_user(request.user, Lead.objects.select_related("book", "author_profile", "brief", "primary_contact")).prefetch_related(
            "evidence", "videos", "contact_candidates__checks", "social_audits"
        ),
        pk=pk,
    )
    field_evidence = {}
    field_evidence_list = {}
    for evidence in lead.evidence.all():
        field_evidence_list.setdefault(evidence.field_name, []).append(evidence)
    for field, evidences in field_evidence_list.items():
        evidences.sort(key=lambda item: item.confidence, reverse=True)
        field_evidence[field] = evidences[0]
    author = lead.author_profile
    contact_rows_raw = [
        ("Official website", author.canonical_website if author else "", "canonical_website"),
        ("Contact page", author.contact_page_url if author else "", "contact_page_url"),
        ("Public email", lead.public_email, "public_email"),
        ("Public phone", lead.public_phone, "public_phone"),
        ("Location", lead.location, "location"),
        ("Instagram", author.instagram_url if author else "", "instagram_url"),
        ("Facebook", author.facebook_url if author else "", "facebook_url"),
        ("TikTok", author.tiktok_url if author else "", "tiktok_url"),
        ("YouTube", author.youtube_url if author else "", "youtube_url"),
        ("LinkedIn", author.linkedin_url if author else "", "linkedin_url"),
        ("Publisher URL", author.publisher_url if author else "", "publisher_url"),
    ]
    contact_rows = []
    contact_fields = {"public_email", "public_phone", "representation_email", "publicist_email"}
    for label, value, field in contact_rows_raw:
        row_evidences = []
        if value:
            row_evidences = [
                item
                for item in field_evidence_list.get(field, [])
                if str(item.field_value) == str(value) and not is_catalog_or_platform_source(item.source_url)
            ][:3]
        if field in contact_fields:
            is_verified = bool(value and has_verified_contact_source(lead, field, value, min_confidence=0.55 if field == "public_phone" else 0.6))
        else:
            is_verified = bool(
                value
                and row_evidences
                and any(
                    item.confidence >= 0.65
                    and item.evidence_type in {"contact_page", "official_author_site", "publisher_site", "manual", "groq_extraction", "social_profile"}
                    for item in row_evidences
                )
            )
        contact_rows.append(
            {
                "label": label,
                "value": value,
                "evidence": row_evidences[0] if row_evidences else field_evidence.get(field),
                "evidences": row_evidences,
                "is_verified": is_verified,
                "is_url": str(value).startswith("http"),
            }
        )

    # =========================================================================
    # ADVANCED DATA CONSOLIDATION & CONSENSUS AUDITING ENGINE
    # =========================================================================
    author_name = lead.book.author_name.strip() if lead.book.author_name else ""
    sibling_leads = []
    sibling_books = []
    
    if author_name:
        # Cross-reference by name or linked AuthorProfile ID
        q_filter_lead = Q(book__author_name__iexact=author_name)
        if lead.author_profile_id:
            q_filter_lead |= Q(author_profile_id=lead.author_profile_id)
            
        sibling_leads = list(
            lead_queryset_for_user(request.user, Lead.objects.filter(q_filter_lead))
            .select_related("book", "author_profile")
            .prefetch_related("evidence")
            .distinct()
        )
        
        q_filter_book = Q(author_name__iexact=author_name)
        if lead.author_profile_id:
            q_filter_book |= Q(leads__author_profile_id=lead.author_profile_id)
            
        sibling_books_query = Book.objects.filter(q_filter_book)
        if not is_management(request.user):
            sibling_books_query = sibling_books_query.filter(
                leads__assignments__assigned_to=request.user,
                leads__assignments__is_current=True,
            )
        sibling_books = list(
            sibling_books_query
            .select_related("research_run")
            .prefetch_related("leads")
            .distinct()
        )
    else:
        sibling_leads = [lead]
        if lead.book:
            sibling_books = [lead.book]

    # Contact channels to check for cross-lead consistency
    fields_to_check = [
        # (field_key, display_label, is_profile_field)
        ("public_email", "Public Email", False),
        ("public_phone", "Public Phone", False),
        ("location", "Location", False),
        ("representation_email", "Representation Email", False),
        ("publicist_email", "Publicist Email", False),
        ("canonical_website", "Official Website", True),
        ("contact_page_url", "Contact Page", True),
        ("instagram_url", "Instagram", True),
        ("facebook_url", "Facebook", True),
        ("tiktok_url", "TikTok", True),
        ("youtube_url", "YouTube", True),
        ("linkedin_url", "LinkedIn", True),
        ("goodreads_url", "Goodreads", True),
        ("amazon_author_url", "Amazon Author Page", True),
        ("publisher_url", "Publisher URL", True),
    ]

    field_observations = {f_key: [] for f_key, _, _ in fields_to_check}

    for s_lead in sibling_leads:
        ap = s_lead.author_profile
        for f_key, _, is_profile in fields_to_check:
            val = ""
            if is_profile:
                if ap:
                    val = getattr(ap, f_key, "")
            else:
                val = getattr(s_lead, f_key, "")
            
            if val:
                val_clean = str(val).strip()
                # Normalize for casing differences
                if "@" in val_clean or val_clean.startswith("http") or f_key in ["public_email", "representation_email", "publicist_email"]:
                    val_clean_norm = val_clean.lower()
                else:
                    val_clean_norm = val_clean
                
                # Retrieve scraping evidence matching this specific value
                matching_evidences = []
                for ev in s_lead.evidence.all():
                    if ev.field_name == f_key and str(ev.field_value).strip().lower() == val_clean_norm:
                        matching_evidences.append({
                            "source_url": ev.source_url,
                            "evidence_type": ev.evidence_type,
                            "confidence": int(ev.confidence * 100) if ev.confidence else 0
                        })
                
                source_info = {
                    "lead_id": str(s_lead.id),
                    "book_title": s_lead.book.title if s_lead.book else "Unknown Book",
                    "book_asin": s_lead.book.asin if s_lead.book else "",
                    "publisher": s_lead.book.publisher or "Self-published",
                    "publication_date": s_lead.book.publication_date or "—",
                    "source_provider": s_lead.book.source_provider or "Unknown Source",
                    "evidences": matching_evidences
                }
                
                field_observations[f_key].append({
                    "norm": val_clean_norm,
                    "orig": val_clean,
                    "source": source_info
                })

    consensus_contacts = {}
    for f_key, f_label, is_profile in fields_to_check:
        obs = field_observations[f_key]
        
        # Get current value for this specific lead
        current_val = ""
        if is_profile:
            if lead.author_profile:
                current_val = getattr(lead.author_profile, f_key, "")
        else:
            current_val = getattr(lead, f_key, "")
        current_val = str(current_val).strip()
        
        # Normalize current value
        if "@" in current_val or current_val.startswith("http") or f_key in ["public_email", "representation_email", "publicist_email"]:
            current_val_norm = current_val.lower()
        else:
            current_val_norm = current_val

        if not obs:
            consensus_contacts[f_key] = {
                "label": f_label,
                "consensus_value": "",
                "confidence_percentage": 0,
                "total_records": 0,
                "status": "missing",  # missing, consistent, conflict, consistent_with_conflict, available_elsewhere
                "current_value": current_val,
                "alternatives": [],
                "is_url": False,
                "consensus_sources": [],
            }
            continue
            
        # Count frequencies
        norm_counts = Counter([item["norm"] for item in obs])
        total_obs = len(obs)
        
        # Get consensus (most common)
        most_common_norm, count = norm_counts.most_common(1)[0]
        most_common_orig = next(item["orig"] for item in obs if item["norm"] == most_common_norm)
        
        confidence_pct = int((count / total_obs) * 100)
        
        # Classify consistency status
        if not current_val:
            status = "available_elsewhere"
        elif current_val_norm != most_common_norm:
            status = "conflict"
        elif len(norm_counts) > 1:
            status = "consistent_with_conflict"
        else:
            status = "consistent"
            
        # Compile sources matching the consensus value
        consensus_sources = [item["source"] for item in obs if item["norm"] == most_common_norm]
            
        # Alternatives (other than consensus)
        alternatives = []
        for norm_val, cnt in norm_counts.items():
            if norm_val != most_common_norm:
                orig_val = next(item["orig"] for item in obs if item["norm"] == norm_val)
                pct = int((cnt / total_obs) * 100)
                alt_sources = [item["source"] for item in obs if item["norm"] == norm_val]
                alternatives.append({
                    "value": orig_val,
                    "count": cnt,
                    "percentage": pct,
                    "sources": alt_sources
                })
                
        consensus_contacts[f_key] = {
            "label": f_label,
            "consensus_value": most_common_orig,
            "confidence_percentage": confidence_pct,
            "total_records": total_obs,
            "matching_records": count,
            "status": status,
            "current_value": current_val,
            "alternatives": alternatives,
            "is_url": most_common_orig.startswith("http"),
            "consensus_sources": consensus_sources,
        }

    # =========================================================================
    # AUTHOR'S BIBLIOGRAPHY / PUBLISHING FOOTPRINT
    # =========================================================================
    portfolio_books = []
    seen_book_asins = set()
    seen_book_titles = set()
    
    # Compile from sibling leads (guarantees we have outreach results)
    for s_lead in sibling_leads:
        bk = s_lead.book
        if not bk:
            continue
        asin_key = (bk.asin or "").upper()
        title_key = bk.title.strip().lower()
        
        if asin_key and asin_key in seen_book_asins:
            continue
        if not asin_key and title_key in seen_book_titles:
            continue
            
        if asin_key:
            seen_book_asins.add(asin_key)
        seen_book_titles.add(title_key)
        
        portfolio_books.append({
            "title": bk.title,
            "cover_image_url": bk.cover_image_url,
            "asin": bk.asin,
            "publisher": bk.publisher or "Self-published",
            "publication_date": bk.publication_date or "—",
            "rating": bk.rating,
            "review_count": bk.review_count,
            "amazon_book_url": bk.amazon_book_url,
            "source_provider": bk.source_provider,
            "marketplace": bk.research_run.marketplace,
            "run_keyword": bk.research_run.keyword or bk.research_run.source_provider,
            "lead_id": s_lead.id,
            "verification_status": s_lead.verification_status,
            "verification_score": s_lead.verification_score,
            "has_lead": True,
            # Sourced contacts specifically for this book
            "email": s_lead.public_email.strip() if s_lead.public_email else "",
            "phone": s_lead.public_phone.strip() if s_lead.public_phone else "",
            "representation_email": s_lead.representation_email.strip() if s_lead.representation_email else "",
            "publicist_email": s_lead.publicist_email.strip() if s_lead.publicist_email else "",
            "location": s_lead.location.strip() if s_lead.location else "",
            "website": s_lead.author_profile.canonical_website.strip() if (s_lead.author_profile and s_lead.author_profile.canonical_website) else "",
        })
        
    # Compile from direct Books (grabs non-lead or incomplete search results)
    for bk in sibling_books:
        asin_key = (bk.asin or "").upper()
        title_key = bk.title.strip().lower()
        
        if asin_key and asin_key in seen_book_asins:
            continue
        if not asin_key and title_key in seen_book_titles:
            continue
            
        if asin_key:
            seen_book_asins.add(asin_key)
        seen_book_titles.add(title_key)
        
        ld = bk.leads.first()
        
        portfolio_books.append({
            "title": bk.title,
            "cover_image_url": bk.cover_image_url,
            "asin": bk.asin,
            "publisher": bk.publisher or "Self-published",
            "publication_date": bk.publication_date or "—",
            "rating": bk.rating,
            "review_count": bk.review_count,
            "amazon_book_url": bk.amazon_book_url,
            "source_provider": bk.source_provider,
            "marketplace": bk.research_run.marketplace,
            "run_keyword": bk.research_run.keyword or bk.research_run.source_provider,
            "lead_id": ld.id if ld else None,
            "verification_status": ld.verification_status if ld else "—",
            "verification_score": ld.verification_score if ld else 0,
            "has_lead": ld is not None,
            # Sourced contacts specifically for this book
            "email": ld.public_email.strip() if (ld and ld.public_email) else "",
            "phone": ld.public_phone.strip() if (ld and ld.public_phone) else "",
            "representation_email": ld.representation_email.strip() if (ld and ld.representation_email) else "",
            "publicist_email": ld.publicist_email.strip() if (ld and ld.publicist_email) else "",
            "location": ld.location.strip() if (ld and ld.location) else "",
            "website": ld.author_profile.canonical_website.strip() if (ld and ld.author_profile and ld.author_profile.canonical_website) else "",
        })

    current_assignment = lead.assignments.filter(is_current=True).select_related("assigned_to", "assigned_by").first()
    return render(
        request,
        "leadfinder/lead_detail.html",
        {
            "lead": lead,
            "field_evidence": field_evidence,
            "contact_rows": contact_rows,
            "consensus_contacts": consensus_contacts,
            "portfolio_books": portfolio_books,
            "sibling_leads_count": len(sibling_leads),
            "current_assignment": current_assignment,
            "assignment_form": LeadAssignmentUpdateForm(instance=current_assignment) if current_assignment else None,
        },
    )


@roles_required("super_admin", "admin")
@require_POST
def run_stop(request, pk):
    run = get_object_or_404(ResearchRun, pk=pk)
    if run.status in {"pending", "running"}:
        run.mark_canceled()
        messages.success(request, "Stop requested. The runner will exit at the next safe checkpoint.")
    else:
        messages.info(request, f"Run is already {run.status}.")
    
    redirect_url = request.POST.get("next") or request.GET.get("next") or request.META.get("HTTP_REFERER")
    if redirect_url and _is_safe_redirect_url(request, redirect_url):
        return redirect(redirect_url)
    return redirect("leadfinder:run_detail", pk=run.id)


@roles_required("super_admin", "admin")
@require_POST
def run_retry(request, pk):
    old_run = get_object_or_404(ResearchRun, pk=pk)
    if old_run.status in {"pending", "running"}:
        messages.warning(request, "This run is still in progress. Stop it before starting a retry.")
        return redirect("leadfinder:run_detail", pk=old_run.id)
    new_settings = dict(old_run.settings_json or {})
    if old_run.source_provider == "ddgs":
        new_settings.setdefault("google_books_fallback", True)
    run = ResearchRun.objects.create(
        keyword=old_run.keyword,
        source_provider=old_run.source_provider,
        marketplace=old_run.marketplace,
        max_books=old_run.max_books,
        settings_json=new_settings,
    )
    _start_research_pipeline(run.id)
    messages.success(request, "Retry started as a new research run.")
    return redirect("leadfinder:run_detail", pk=run.id)


@require_POST
def lead_action(request, pk, action):
    lead = get_object_or_404(lead_queryset_for_user(request.user), pk=pk)
    if not is_management(request.user) and action not in {"add-note"}:
        messages.error(request, "Sales users update leads through the assignment task panel.")
        return redirect("leadfinder:lead_detail", pk=lead.id)
    if action == "delete":
        lead.delete()
        messages.success(request, "Lead deleted successfully.")
        return redirect("leadfinder:lead_list")
    elif action == "approve":
        ok, errors = can_approve_lead(lead, manual_identity_override=True)
        if ok:
            lead.manual_review_status = "approved"
            messages.success(request, "Lead approved.")
        else:
            messages.error(request, "Cannot approve: " + "; ".join(errors))
    elif action == "reject":
        lead.manual_review_status = "rejected"
        messages.success(request, "Lead rejected.")
    elif action == "do-not-contact":
        lead.manual_review_status = "do_not_contact"
        lead.do_not_contact = True
        messages.success(request, "Lead marked do-not-contact.")
    elif action == "reverify":
        harvest_social_profiles(lead)
        verify_lead_contacts(lead, check_network=True)
        lead.refresh_from_db()
        messages.success(
            request,
            f"Contact verification completed: {lead.get_verification_status_display()} "
            f"({lead.verification_score}/100).",
        )
    elif action == "add-note":
        if not request.POST.get("note", "").strip():
            messages.warning(request, "Enter a note before saving the activity log.")
            return redirect("leadfinder:lead_detail", pk=lead.id)
        messages.success(request, "Activity note saved without changing review status.")
    else:
        messages.error(request, "Unknown lead action.")
        return redirect("leadfinder:lead_detail", pk=lead.id)
    note = request.POST.get("note", "").strip()
    if note:
        lead.notes = f"{lead.notes}\n{note}".strip()
    lead.save()
    return redirect("leadfinder:lead_detail", pk=lead.id)


@require_POST
def lead_bulk_action(request):
    lead_ids = request.POST.getlist("lead_ids")
    action = request.POST.get("action")
    redirect_url = request.POST.get("next") or request.META.get("HTTP_REFERER") or "leadfinder:lead_list"
    if not _is_safe_redirect_url(request, redirect_url):
        redirect_url = "leadfinder:lead_list"
    
    if not lead_ids:
        messages.warning(request, "No leads were selected.")
        return redirect(redirect_url)
        
    accessible_leads = lead_queryset_for_user(request.user).filter(id__in=lead_ids)
    if action == "bulk_assign" and is_management(request.user):
        salesperson = get_object_or_404(
            get_user_model().objects.filter(is_active=True, leadfinder_profile__role="sales"),
            pk=request.POST.get("assigned_to"),
        )
        assigned = assign_leads(
            lead_ids=accessible_leads.values_list("id", flat=True),
            salesperson=salesperson,
            assigned_by=request.user,
            replace_current=True,
        )
        messages.success(request, f"Assigned {assigned} unique leads to {salesperson.get_full_name() or salesperson.username}.")
    elif action == "bulk_unassign" and is_management(request.user):
        updated = LeadAssignment.objects.filter(lead__in=accessible_leads, is_current=True).update(is_current=False)
        messages.success(request, f"Returned {updated} leads to the unassigned pool.")
    elif action == "bulk_reject" and is_management(request.user):
        updated = accessible_leads.update(manual_review_status="rejected")
        messages.success(request, f"Successfully rejected {updated} leads.")
    elif action == "bulk_delete" and is_management(request.user):
        deleted_count, _ = accessible_leads.delete()
        messages.success(request, f"Successfully deleted {deleted_count} leads.")
    else:
        messages.error(request, "Invalid bulk action requested.")
        
    return redirect(redirect_url)


def export_leads_csv(request):
    leads = lead_queryset_for_user(request.user).exclude(do_not_contact=True).order_by("-verification_score", "-created_at")
    valid_only = request.GET.get("valid_only", "0").lower() not in {"0", "false", "all", "no", ""}
    if valid_only:
        leads = EligibilityPolicy.verified_ready(leads)
    limit = request.GET.get("limit")
    if limit:
        try:
            leads = leads[: max(1, min(int(limit), 700))]
        except ValueError:
            pass
    return export_leads_response(leads)


def export_leads_xlsx_view(request):
    leads = lead_queryset_for_user(request.user).exclude(do_not_contact=True).order_by("-verification_score", "-created_at")
    valid_only = request.GET.get("valid_only", "0").lower() not in {"0", "false", "all", "no", ""}
    if valid_only:
        leads = EligibilityPolicy.verified_ready(leads)
    limit = request.GET.get("limit")
    if limit:
        try:
            leads = leads[: max(1, min(int(limit), 700))]
        except ValueError:
            pass
    return export_leads_xlsx(leads)


@require_POST
def lead_assignment_update(request, pk):
    assignment = get_object_or_404(
        LeadAssignment.objects.select_related("lead", "assigned_to"),
        pk=pk,
        is_current=True,
    )
    ensure_lead_access(request.user, assignment.lead)
    if user_role(request.user) == "sales" and assignment.assigned_to_id != request.user.id:
        return HttpResponse("This task is not assigned to you.", status=403)
    form = LeadAssignmentUpdateForm(request.POST, instance=assignment)
    if form.is_valid():
        previous_status = assignment.status
        assignment = form.save(commit=False)
        now = timezone.now()
        if assignment.status != "pending" and not assignment.started_at:
            assignment.started_at = now
        if assignment.status in {"completed", "converted", "not_interested", "no_response", "invalid_data", "duplicate"}:
            assignment.completed_at = assignment.completed_at or now
        elif previous_status != assignment.status:
            assignment.completed_at = None
        if assignment.status == "converted":
            assignment.converted_at = assignment.converted_at or now
        elif previous_status == "converted":
            assignment.converted_at = None
        assignment.save()
        messages.success(request, "Lead task updated.")
    else:
        messages.error(request, "Please correct the task status fields.")
    return redirect("leadfinder:lead_detail", pk=assignment.lead_id)


@roles_required("super_admin", "admin")
def assignment_schedule_list(request):
    schedules = LeadAssignmentSchedule.objects.select_related("salesperson", "created_by").prefetch_related("runs")
    return render(
        request,
        "leadfinder/assignment_schedule_list.html",
        {
            "schedules": schedules,
            "recent_runs": LeadAssignmentScheduleRun.objects.select_related("schedule", "schedule__salesperson")[:12],
            "stats": {
                "active": schedules.filter(is_active=True).count(),
                "daily_capacity": sum(item.daily_lead_count for item in schedules.filter(is_active=True)),
                "assigned": LeadAssignment.objects.filter(is_current=True).count(),
                "unassigned": Lead.objects.filter(do_not_contact=False).exclude(assignments__is_current=True).count(),
            },
        },
    )


@roles_required("super_admin", "admin")
def assignment_schedule_create(request):
    if request.method == "POST":
        form = LeadAssignmentScheduleForm(request.POST)
        if form.is_valid():
            schedule = form.save(commit=False)
            schedule.created_by = request.user
            schedule.next_run_at = calculate_next_assignment_at(schedule)
            schedule.save()
            messages.success(request, f'Assignment schedule “{schedule.name}” created.')
            return redirect("leadfinder:assignment_schedule_list")
    else:
        form = LeadAssignmentScheduleForm()
    return render(request, "leadfinder/assignment_schedule_form.html", {"form": form, "page_title": "Create assignment schedule"})


@roles_required("super_admin", "admin")
def assignment_schedule_edit(request, pk):
    schedule = get_object_or_404(LeadAssignmentSchedule, pk=pk)
    if request.method == "POST":
        form = LeadAssignmentScheduleForm(request.POST, instance=schedule)
        if form.is_valid():
            schedule = form.save(commit=False)
            schedule.next_run_at = calculate_next_assignment_at(schedule)
            schedule.save()
            messages.success(request, f'Assignment schedule “{schedule.name}” updated.')
            return redirect("leadfinder:assignment_schedule_list")
    else:
        form = LeadAssignmentScheduleForm(instance=schedule)
    return render(request, "leadfinder/assignment_schedule_form.html", {"form": form, "page_title": "Edit assignment schedule", "schedule": schedule})


@roles_required("super_admin", "admin")
@require_POST
def assignment_schedule_toggle(request, pk):
    schedule = get_object_or_404(LeadAssignmentSchedule, pk=pk)
    schedule.is_active = not schedule.is_active
    if schedule.is_active:
        schedule.next_run_at = calculate_next_assignment_at(schedule)
    schedule.save(update_fields=["is_active", "next_run_at", "updated_at"])
    messages.success(request, f'“{schedule.name}” is now {"active" if schedule.is_active else "paused"}.')
    return redirect("leadfinder:assignment_schedule_list")


@roles_required("super_admin", "admin")
@require_POST
def assignment_schedule_run(request, pk):
    run = execute_assignment_schedule(pk, force=True)
    messages.success(request, run.message)
    return redirect("leadfinder:assignment_schedule_list")


@roles_required("super_admin", "admin")
@require_POST
def assignment_schedule_delete(request, pk):
    schedule = get_object_or_404(LeadAssignmentSchedule, pk=pk)
    name = schedule.name
    schedule.delete()
    messages.success(request, f'“{name}” was deleted. Existing assignments were preserved.')
    return redirect("leadfinder:assignment_schedule_list")


@roles_required("super_admin", "admin")
def team_list(request):
    users = get_user_model().objects.select_related("leadfinder_profile").annotate(
        assigned_count=Count("lead_assignments", filter=Q(lead_assignments__is_current=True)),
        completed_count=Count("lead_assignments", filter=Q(lead_assignments__is_current=True, lead_assignments__status="completed")),
        converted_count=Count("lead_assignments", filter=Q(lead_assignments__is_current=True, lead_assignments__status="converted")),
    ).order_by("first_name", "last_name", "username")
    current_role = user_role(request.user)
    form = TeamMemberCreateForm(request.POST or None, force_sales_role=current_role == "admin")
    if request.method == "POST" and form.is_valid():
        user = form.save()
        messages.success(request, f"{user.get_full_name() or user.username} can now sign in.")
        return redirect("leadfinder:team_list")
    return render(request, "leadfinder/team_list.html", {"users": users, "form": form, "can_choose_roles": current_role == "super_admin"})


@roles_required("super_admin", "admin")
@require_POST
def team_member_toggle(request, pk):
    user = get_object_or_404(get_user_model(), pk=pk)
    if user == request.user or user.is_superuser or (user_role(request.user) == "admin" and user_role(user) != "sales"):
        messages.error(request, "You cannot change that account.")
    else:
        user.is_active = not user.is_active
        user.save(update_fields=["is_active"])
        messages.success(request, f"Account {user.username} is now {'active' if user.is_active else 'disabled'}.")
    return redirect("leadfinder:team_list")


@roles_required("super_admin")
@require_POST
def team_member_role_update(request, pk):
    user = get_object_or_404(get_user_model(), pk=pk)
    role = request.POST.get("role", "")
    if user == request.user:
        messages.error(request, "You cannot change your own role.")
    elif role not in dict(UserProfile.ROLE_CHOICES):
        messages.error(request, "Choose a valid role.")
    else:
        profile, _ = UserProfile.objects.get_or_create(user=user)
        profile.role = role
        profile.is_available_for_assignment = role == "sales" and request.POST.get("available") == "on"
        profile.save(update_fields=["role", "is_available_for_assignment", "updated_at"])
        user.is_superuser = role == "super_admin"
        user.is_staff = role == "super_admin"
        user.save(update_fields=["is_superuser", "is_staff"])
        messages.success(request, f"{user.username} now has the {profile.get_role_display()} role.")
    return redirect("leadfinder:team_list")


_env_file_lock = threading.Lock()


def save_env_settings(updates: dict[str, str]) -> None:
    """Persist approved settings without exposing or corrupting secret values.

    Note: os.environ changes apply to this process only; the scheduler loop
    process picks up new keys on its next restart.
    """
    env_path = os.path.join(settings.BASE_DIR, ".env")
    for k, v in updates.items():
        if "\n" in v or "\r" in v:
            raise ValueError(f"{k} must be a single-line value.")
        if v:
            os.environ[k] = v
        else:
            os.environ.pop(k, None)

    with _env_file_lock:
        lines = []
        if os.path.exists(env_path):
            with open(env_path, "r", encoding="utf-8") as f:
                lines = f.readlines()

        updated_keys = set()
        new_lines = []

        for line in lines:
            stripped = line.strip()
            if not stripped or stripped.startswith("#"):
                new_lines.append(line)
                continue

            if "=" in stripped:
                k = stripped.split("=", 1)[0].strip()
                if k in updates:
                    escaped_value = updates[k].replace("\\", "\\\\").replace('"', '\\"')
                    new_lines.append(f'{k}="{escaped_value}"\n')
                    updated_keys.add(k)
                else:
                    new_lines.append(line)
            else:
                new_lines.append(line)

        for k, v in updates.items():
            if k not in updated_keys:
                escaped_value = v.replace("\\", "\\\\").replace('"', '\\"')
                new_lines.append(f'{k}="{escaped_value}"\n')

        # Write via a temp file + atomic replace so a crash never truncates .env.
        tmp_path = env_path + ".tmp"
        with open(tmp_path, "w", encoding="utf-8") as f:
            f.writelines(new_lines)
        os.replace(tmp_path, env_path)


@roles_required("super_admin", "admin")
def settings_help(request):
    api_key_names = [
        "SEARCH_PROVIDER",
        "GROQ_API_KEY",
        "GROQ_MODEL",
        "TAVILY_API_KEY",
        "GOOGLE_API_KEY",
        "GOOGLE_CSE_ID",
        "BRAVE_API_KEY",
        "YOUTUBE_API_KEY",
        "GOOGLE_BOOKS_API_KEY",
        "AMAZON_CREATORS_CLIENT_ID",
        "AMAZON_CREATORS_CLIENT_SECRET",
    ]
    secret_names = {
        "GROQ_API_KEY",
        "TAVILY_API_KEY",
        "GOOGLE_API_KEY",
        "GOOGLE_CSE_ID",
        "BRAVE_API_KEY",
        "YOUTUBE_API_KEY",
        "GOOGLE_BOOKS_API_KEY",
        "AMAZON_CREATORS_CLIENT_ID",
        "AMAZON_CREATORS_CLIENT_SECRET",
    }

    if request.method == "POST":
        action = request.POST.get("settings_action", "save")
        remove_key = request.POST.get("remove_key", "")
        if action == "remove":
            if remove_key not in secret_names:
                messages.error(request, "Choose a configured secret to remove.")
            else:
                save_env_settings({remove_key: ""})
                messages.success(request, f"{remove_key} was removed from local configuration.")
        elif action == "test":
            configured_count = sum(bool(os.getenv(key)) for key in secret_names)
            provider = request.POST.get("SEARCH_PROVIDER", os.getenv("SEARCH_PROVIDER", "ddgs"))
            dependencies = {
                "tavily": ("TAVILY_API_KEY",),
                "google": ("GOOGLE_API_KEY", "GOOGLE_CSE_ID"),
                "brave": ("BRAVE_API_KEY",),
            }
            missing = [key for key in dependencies.get(provider, ()) if not os.getenv(key)]
            if missing:
                messages.warning(request, f"{provider.title()} is not ready: add {', '.join(missing)}.")
            else:
                messages.success(
                    request,
                    f"Configuration check passed: {configured_count} stored credential(s). "
                    "No provider request was sent.",
                )
        else:
            updates = {}
            for key in api_key_names:
                value = request.POST.get(key, "").strip()
                if key not in secret_names or value:
                    updates[key] = value
            try:
                save_env_settings(updates)
            except ValueError as exc:
                messages.error(request, str(exc))
            else:
                messages.success(request, "Provider settings updated. Blank secret fields were left unchanged.")
        return redirect("leadfinder:settings_help")

    def configured(name: str) -> bool:
        return bool(os.getenv(name))

    def installed(module_name: str) -> bool:
        return importlib.util.find_spec(module_name) is not None

    booklife_direct_fetch = os.getenv("BOOKLIFE_DIRECT_FETCH_ALLOWED", "").lower() in {"1", "true", "yes", "on"}

    # Only non-sensitive configuration is rendered into HTML.  Secret values
    # are represented by a configured/missing boolean and can only be replaced.
    current_env = {key: os.getenv(key, "") for key in ("SEARCH_PROVIDER", "GROQ_MODEL")}

    context = {
        "search_provider": os.getenv("SEARCH_PROVIDER", "ddgs"),
        "current_env": current_env,
        "keys": {
            "GOOGLE_API_KEY": configured("GOOGLE_API_KEY"),
            "GOOGLE_CSE_ID": configured("GOOGLE_CSE_ID"),
            "TAVILY_API_KEY": configured("TAVILY_API_KEY"),
            "BRAVE_API_KEY": configured("BRAVE_API_KEY"),
            "GROQ_API_KEY": configured("GROQ_API_KEY"),
            "YOUTUBE_API_KEY": configured("YOUTUBE_API_KEY"),
            "GOOGLE_BOOKS_API_KEY": configured("GOOGLE_BOOKS_API_KEY"),
            "AMAZON_CREATORS_CLIENT_ID": configured("AMAZON_CREATORS_CLIENT_ID"),
            "AMAZON_CREATORS_CLIENT_SECRET": configured("AMAZON_CREATORS_CLIENT_SECRET"),
        },
        "scraping_tools": [
            {
                "name": "DDGS",
                "purpose": "Default no-key public web search fallback.",
                "status": "Available" if installed("ddgs") else "Missing package",
            },
            {
                "name": "Google Custom Search",
                "purpose": "Optional Google search provider for premium discovery results.",
                "status": "Configured" if (configured("GOOGLE_API_KEY") and configured("GOOGLE_CSE_ID")) else "Missing API key/CSE ID",
            },
            {
                "name": "Tavily",
                "purpose": "Optional API-backed web search provider.",
                "status": "Configured" if configured("TAVILY_API_KEY") else "Missing API key",
            },
            {
                "name": "Brave Search",
                "purpose": "Optional API-backed web search provider.",
                "status": "Configured" if configured("BRAVE_API_KEY") else "Missing API key",
            },
            {
                "name": "Google Books",
                "purpose": "Structured book metadata provider for cleaner title, author, ISBN, publisher, category, rating, and cover data.",
                "status": "API key configured" if configured("GOOGLE_BOOKS_API_KEY") else "No-key public mode",
            },
            {
                "name": "ISBN intelligence",
                "purpose": "Checksum agreement, ISBN-10/13 conversion, exact-source reconciliation, and SVG barcode output.",
                "status": "Available" if all(installed(name) for name in ("isbnlib", "pyisbn", "barcode")) else "Missing package",
            },
            {
                "name": "BookLife",
                "purpose": "Category-based project discovery that feeds author/contact enrichment.",
                "status": "Direct fetch enabled" if booklife_direct_fetch else "Robots-respecting mode",
            },
            {
                "name": "Scrapy",
                "purpose": "Batch author-site scraping commands.",
                "status": "Available" if installed("scrapy") else "Missing package",
            },
            {
                "name": "Groq",
                "purpose": "Optional structured AI extraction and briefs.",
                "status": "Configured" if configured("GROQ_API_KEY") else "Fallback mode",
            },
        ],
        "local_checks": [
            {
                "label": "Run app tests",
                "command": r"..\.venv\Scripts\python.exe -m pytest",
            },
            {
                "label": "Run Django checks",
                "command": r"..\.venv\Scripts\python.exe manage.py check",
            },
            {
                "label": "One-command local verification",
                "command": r".\scripts\test_leadfinder.ps1",
            },
        ],
    }
    return render(request, "leadfinder/settings_help.html", context)


@roles_required("super_admin", "admin")
def isbn_search(request):
    run_id = request.GET.get("run_id")
    active_run = None
    if run_id:
        try:
            active_run = ResearchRun.objects.get(id=run_id)
        except (ResearchRun.DoesNotExist, ValueError):
            pass

    if request.method == "POST":
        form = ISBNSearchForm(request.POST)
        if form.is_valid():
            identifier_analyses = form.identifier_analyses
            isbn_list = [item.canonical for item in identifier_analyses]

            if not isbn_list:
                messages.error(request, "Please enter at least one valid ASIN or ISBN.")
                return redirect("leadfinder:isbn_search")

            run_video = form.cleaned_data["run_video_search"]
            run_ai = form.cleaned_data["run_groq_ai_extraction"]

            # Try to see if this book already exists in DB with valid author
            if len(isbn_list) == 1:
                analysis = identifier_analyses[0]
                aliases = [value for value in {analysis.canonical, analysis.isbn10, analysis.isbn13} if value]
                book = Book.objects.filter(asin__in=aliases).exclude(author_name="").order_by("-book_data_confidence").first()
                if (
                    book
                    and analysis.identifier_type in {"isbn10", "isbn13"}
                    and not (book.source_raw_json or {}).get("metadata_resolution")
                ):
                    book = None
                if book:
                    lead = Lead.objects.filter(book=book).select_related("book", "author_profile").first()
                    if lead:
                        messages.info(request, f"Book details and lead for {analysis.display} retrieved from the evidence cache.")
                        return redirect("leadfinder:lead_detail", pk=lead.id)

            # Keyword description:
            if len(isbn_list) == 1:
                keyword_text = f"ISBN Search: {isbn_list[0]}"
            else:
                keyword_text = f"Batch ISBN Search ({len(isbn_list)} items)"

            # Create a Manual/ISBN ResearchRun
            run = ResearchRun.objects.create(
                keyword=keyword_text,
                source_provider="manual",
                max_books=len(isbn_list),
                settings_json={
                    "require_amazon_url": False,
                    "require_public_email": False,
                    "include_social_only_leads": True,
                    "run_video_search": run_video,
                    "run_groq_ai_extraction": run_ai,
                }
            )

            # Create initial Book records
            for analysis in identifier_analyses:
                code = analysis.canonical
                # A typed ISBN/ASIN is not proof that an Amazon product exists.
                # Preserve an actual user-supplied Amazon product URL only; the
                # pipeline may later attach independently found product evidence.
                supplied_amazon_url = (
                    normalize_amazon_book_url(
                        analysis.raw, os.getenv("AMAZON_ASSOCIATE_TAG") or None
                    )
                    if is_amazon_url(analysis.raw) and extract_asin(analysis.raw)
                    else ""
                )
                Book.objects.create(
                    research_run=run,
                    title=f"Book for {analysis.identifier_type.upper()} {code}",
                    asin=code,
                    amazon_book_url=supplied_amazon_url,
                    amazon_source_url=supplied_amazon_url,
                    normalized_key=normalized_book_key(f"Book for {analysis.identifier_type.upper()} {code}", "", code),
                    book_data_confidence=analysis.confidence,
                    source_provider="manual",
                    source_raw_json={
                        "processing_status": "pending",
                        "identifier_intelligence": analysis.as_dict(),
                    },
                )

            # Start background pipeline
            _start_research_pipeline(run.id)

            return redirect(f"{reverse('leadfinder:isbn_search')}?run_id={run.id}")
    else:
        form = ISBNSearchForm()

    return render(request, "leadfinder/isbn_search.html", {
        "form": form,
        "active_run": active_run,
        "isbn_stack": ["isbnlib", "pyisbn", "Google Books", "Open Library", "Library of Congress", "Crossref", "Internet Archive", "python-barcode"],
    })


@roles_required("super_admin", "admin")
@require_GET
def isbn_analyze(request):
    analysis = analyze_identifier(request.GET.get("identifier", ""))
    payload = analysis.as_dict()
    if analysis.barcode_available:
        payload["barcode_url"] = reverse("leadfinder:isbn_barcode", kwargs={"identifier": analysis.canonical})
    return JsonResponse(payload, status=200 if analysis.valid else 422)


@roles_required("super_admin", "admin")
@require_GET
def isbn_barcode(request, identifier):
    try:
        svg = generate_barcode_svg(identifier)
    except ValueError as exc:
        return JsonResponse({"error": str(exc)}, status=422)
    response = HttpResponse(svg, content_type="image/svg+xml")
    response["Content-Disposition"] = f'inline; filename="isbn-{analyze_identifier(identifier).isbn13}.svg"'
    response["Cache-Control"] = "public, max-age=86400"
    response["X-Content-Type-Options"] = "nosniff"
    return response


@roles_required("super_admin", "admin")
def isbn_search_status(request, run_id):
    run = get_object_or_404(ResearchRun, id=run_id)
    books = run.books.all().order_by("created_at")

    total_books = books.count()
    completed_books = 0
    failed_books = 0
    pending_books = 0
    processing_books = 0

    tasks_list = []

    for book in books:
        raw = book.source_raw_json or {}
        status = raw.get("processing_status", "pending")

        # Fallback check: if book has processing status as pending/processing but the run itself is finished/failed/canceled,
        # show it as failed or skipped to avoid endless loading UI if the runner crashed.
        if status in ("pending", "processing") and run.status in ("completed", "failed", "canceled"):
            status = "failed"

        if status == "completed":
            completed_books += 1
        elif status == "failed":
            failed_books += 1
        elif status == "processing":
            processing_books += 1
        else:
            pending_books += 1

        # Find lead ID if completed
        lead_id = None
        if status == "completed":
            lead = book.leads.first()
            if lead:
                lead_id = str(lead.id)

        errors = raw.get("book_processing_errors", [])

        tasks_list.append({
            "asin": book.asin,
            "title": book.title,
            "author": book.author_name or "Searching...",
            "status": status,
            "stage": raw.get("processing_stage", ""),
            "detail": raw.get("processing_detail", ""),
            "lead_id": lead_id,
            "errors": errors,
            "identifier": raw.get("identifier_intelligence", {}),
        })

    # Calculate percentage
    finished_count = completed_books + failed_books
    if total_books > 0:
        percentage = int((finished_count / total_books) * 100)
    else:
        percentage = 0

    # Calculate ETC
    etc_str = "Estimating..."
    if run.status == "completed":
        etc_str = "Complete"
    elif run.status in ("failed", "canceled"):
        etc_str = run.status.capitalize()
    else:
        # Active run
        from django.utils import timezone
        started_at = run.started_at or run.created_at
        elapsed = (timezone.now() - started_at).total_seconds()
        if elapsed <= 0:
            elapsed = 1.0
        if finished_count > 0:
            seconds_per_book = elapsed / finished_count
            remaining_books = total_books - finished_count
            remaining_seconds = remaining_books * seconds_per_book
            if remaining_seconds < 60:
                etc_str = f"{int(remaining_seconds)}s remaining"
            else:
                minutes = int(remaining_seconds // 60)
                seconds = int(remaining_seconds % 60)
                etc_str = f"{minutes}m {seconds}s remaining"
        else:
            etc_str = "Estimating..."

    return JsonResponse({
        "run_id": str(run.id),
        "status": run.status,
        "total": total_books,
        "completed": completed_books,
        "processing": processing_books,
        "pending": pending_books,
        "failed": failed_books,
        "percentage": percentage,
        "etc": etc_str,
        "tasks": tasks_list,
    })


@roles_required("super_admin", "admin")
def isbn_lookup(request):
    keyword = request.GET.get("keyword", "").strip()
    pub_year_start = request.GET.get("pub_year_start", "").strip()
    pub_year_end = request.GET.get("pub_year_end", "").strip()
    only_new = request.GET.get("only_new", "false").lower() == "true"
    try:
        max_results = int(request.GET.get("max_results", "50"))
        max_results = max(1, min(max_results, 1000))
    except ValueError:
        max_results = 50
    if not keyword:
        return JsonResponse({"error": "Keyword is required."}, status=400)

    def event_stream():
        import json

        from leadfinder.services.books.isbn_keyword_lookup import stream_keyword_lookup

        for event in stream_keyword_lookup(keyword, pub_year_start, pub_year_end, max_results, only_new=only_new):
            yield json.dumps(event, ensure_ascii=True) + "\n"

    from django.http import StreamingHttpResponse

    response = StreamingHttpResponse(event_stream(), content_type="application/x-ndjson")
    response["X-Accel-Buffering"] = "no"
    response["Cache-Control"] = "no-cache"
    return response
